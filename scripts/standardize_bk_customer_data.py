"""Create a canonical BK workbook from a current customer workbook.

The customer workbook is the data authority.  The configured Output workbook is
used only as the canonical month-by-month header and column-format authority.
"""

from __future__ import annotations

import argparse
import copy
import hashlib
import json
import os
import sys
import tempfile
import zipfile
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

PROJECT_ROOT = Path(__file__).resolve().parents[1]
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

import app.services.excel.payment_sync as payment_sync
from app.rpa_expense.service import (
    RpaExpenseService,
    STATUS_HEADER,
    SUMMARY_HEADERS as RPA_SUMMARY_HEADERS,
)
from app.services.excel.bang_ke import ensure_bang_ke_fee_columns
from app.services.excel.carrier import DETAIL_HEADERS, ensure_bk_detail_sheet
from app.services.excel.daily_sync import (
    SOURCE_HEADER_ALIASES,
    SYNC_FIELDS,
)
from app.services.excel.headers import HeaderResolver, normalize_header
from app.services.excel.posting import (
    FEE_HEADER_ALIASES,
    ExpensePostingService,
)
from app.services.excel.resolvers import MonthSheetService
from app.services.excel.workbook import WorkbookGateway


DATE_NUMBER_FORMAT = "dd/mm/yyyy hh:mm:ss"


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _fast_data_rows(
    worksheet: Any,
    *,
    sqt_column: int,
    container_column: int,
    header_row: int = 1,
) -> list[int]:
    """Equivalent to payment_sync._data_rows without scanning styled tail rows."""

    candidates = sorted(
        {
            row
            for (row, column), cell in worksheet._cells.items()
            if row > header_row
            and column in {sqt_column, container_column}
            and cell.value not in (None, "")
        }
    )
    return [
        row
        for row in candidates
        if payment_sync._parse_sqt(worksheet.cell(row, sqt_column).value)
        is not None
        and payment_sync._container_key(
            worksheet.cell(row, container_column).value
        )
    ]


def _record_rows(worksheet: Any) -> list[int]:
    return _fast_data_rows(
        worksheet,
        sqt_column=1,
        container_column=3,
        header_row=1,
    )


def _record_identity(worksheet: Any, row: int) -> tuple[Any, str]:
    return (
        payment_sync._parse_sqt(worksheet.cell(row, 1).value),
        payment_sync._container_key(worksheet.cell(row, 3).value),
    )


def _column_has_values(worksheet: Any, column: int) -> bool:
    return any(
        row > 1 and cell.value not in (None, "")
        for (row, candidate), cell in worksheet._cells.items()
        if candidate == column
    )


def _trim_styled_tail(worksheet: Any) -> int:
    """Drop empty style-only cells below the last real workbook row."""

    value_rows = [
        row
        for (row, _column), cell in worksheet._cells.items()
        if cell.value not in (None, "")
        or cell.comment is not None
        or cell.hyperlink is not None
    ]
    last_row = max(value_rows, default=1)
    before = len(worksheet._cells)
    worksheet._cells = {
        key: cell
        for key, cell in worksheet._cells.items()
        if key[0] <= last_row
    }
    for row in list(worksheet.row_dimensions):
        if row > last_row:
            del worksheet.row_dimensions[row]
    return before - len(worksheet._cells)


def _copy_column_format(template: Any, target: Any, column: int) -> None:
    letter = get_column_letter(column)
    source = template.column_dimensions[letter]
    destination = target.column_dimensions[letter]
    for attribute in (
        "width",
        "hidden",
        "bestFit",
        "outlineLevel",
        "collapsed",
    ):
        setattr(destination, attribute, copy.copy(getattr(source, attribute)))


def _copy_header_format(template: Any, target: Any, last_column: int) -> None:
    for column in range(1, last_column + 1):
        source = template.cell(1, column)
        destination = target.cell(1, column)
        destination._style = copy.copy(source._style)
        destination.font = copy.copy(source.font)
        destination.fill = copy.copy(source.fill)
        destination.border = copy.copy(source.border)
        destination.alignment = copy.copy(source.alignment)
        destination.protection = copy.copy(source.protection)
        destination.number_format = source.number_format
        _copy_column_format(template, target, column)
    target.row_dimensions[1].height = copy.copy(
        template.row_dimensions[1].height
    )


def _align_business_headers(worksheet: Any, template: Any) -> tuple[int, list[int]]:
    """Remove only proven-empty legacy spacer columns before the summary."""

    target_summary = payment_sync.find_summary_start(template)
    current_summary = payment_sync.find_summary_start(worksheet)
    if target_summary is None or current_summary is None:
        raise RuntimeError(
            f"Missing summary block while aligning {worksheet.title!r}."
        )

    # Column I already contains customer data.  Only its legacy header is blank.
    if normalize_header(worksheet.cell(1, 9).value) == "":
        worksheet.cell(1, 9).value = template.cell(1, 9).value

    deleted: list[int] = []
    current_column = 1
    target_column = 1
    while current_column < current_summary and target_column < target_summary:
        current = normalize_header(worksheet.cell(1, current_column).value)
        expected = normalize_header(template.cell(1, target_column).value)
        if current == expected:
            current_column += 1
            target_column += 1
            continue
        if current == "" and not _column_has_values(worksheet, current_column):
            payment_sync._delete_column(
                worksheet,
                current_column,
                replacement=max(1, current_column - 1),
            )
            deleted.append(current_column)
            current_summary -= 1
            continue
        raise RuntimeError(
            f"Unsafe header mismatch in {worksheet.title!r}: "
            f"column {current_column}={worksheet.cell(1, current_column).value!r}, "
            f"expected {template.cell(1, target_column).value!r}."
        )

    if current_column != current_summary or target_column != target_summary:
        raise RuntimeError(
            f"Business-column count mismatch in {worksheet.title!r}: "
            f"current summary={current_summary}, target summary={target_summary}."
        )
    if payment_sync.find_summary_start(worksheet) != target_summary:
        raise RuntimeError(
            f"Summary position mismatch remains in {worksheet.title!r}."
        )
    return target_summary, deleted


def _canonicalize_sheet(worksheet: Any, template: Any) -> dict[str, Any]:
    ensure_bang_ke_fee_columns(worksheet, header_row=1)
    summary_start, deleted = _align_business_headers(worksheet, template)

    template_summary = payment_sync.find_summary_start(template)
    if summary_start != template_summary:
        raise RuntimeError(f"Unexpected summary position in {worksheet.title!r}.")

    status_column = summary_start + len(RPA_SUMMARY_HEADERS) + 1
    for column in range(1, status_column + 1):
        worksheet.cell(1, column).value = template.cell(1, column).value

    # Refresh after all structure changes, then restore the required combined label.
    payment_sync.refresh_bk_summary_formulas(worksheet)
    worksheet.cell(1, summary_start + 9).value = RPA_SUMMARY_HEADERS[-1]
    worksheet.cell(1, summary_start + 10).value = template.cell(
        1, summary_start + 10
    ).value
    worksheet.cell(1, status_column).value = template.cell(
        1, status_column
    ).value

    _copy_header_format(template, worksheet, status_column)
    rows = _record_rows(worksheet)
    for row in rows:
        worksheet.cell(row, summary_start + 10).number_format = DATE_NUMBER_FORMAT
        worksheet.cell(row, status_column).number_format = "General"

    last_row = max(rows, default=1)
    worksheet.auto_filter.ref = (
        f"A1:{get_column_letter(status_column)}{last_row}"
    )
    return {
        "records": len(rows),
        "first_row": min(rows, default=None),
        "last_row": max(rows, default=None),
        "summary_start": get_column_letter(summary_start),
        "status_column": get_column_letter(status_column),
        "deleted_empty_columns": [get_column_letter(value) for value in deleted],
    }


def _validate_formula_block(worksheet: Any, rows: list[int]) -> None:
    summary_start = payment_sync.find_summary_start(worksheet)
    if summary_start is None:
        raise RuntimeError(f"Missing summary block in {worksheet.title!r}.")
    headers = payment_sync._header_values(worksheet)
    raw = {
        field: payment_sync._required_column(
            headers,
            aliases,
            field,
            before=summary_start,
        )
        for field, aliases in payment_sync.BK_HEADER_ALIASES.items()
        if field
        in {
            "sqt",
            "north_freight",
            "empty_lift",
            "loaded_drop",
            "sea_freight",
            "loaded_lift",
            "empty_drop",
            "vs_do",
            "command_fee",
            "south_freight",
            "storage",
            "repair",
            "overweight",
        }
    }
    letter = get_column_letter
    for row in rows:
        expected = (
            f"={letter(raw['sqt'])}{row}",
            f"={letter(raw['north_freight'])}{row}",
            f"={letter(raw['empty_lift'])}{row}+{letter(raw['loaded_drop'])}{row}",
            f"={letter(raw['sea_freight'])}{row}",
            f"={letter(raw['loaded_lift'])}{row}+{letter(raw['empty_drop'])}{row}+"
            f"{letter(raw['vs_do'])}{row}+{letter(raw['command_fee'])}{row}",
            f"={letter(raw['south_freight'])}{row}",
            f"={letter(raw['storage'])}{row}",
            f"={letter(raw['repair'])}{row}",
            f"={letter(raw['overweight'])}{row}",
            f"={letter(summary_start + 6)}{row}+{letter(summary_start + 8)}{row}",
        )
        actual = tuple(
            worksheet.cell(row, summary_start + offset).value
            for offset in range(10)
        )
        if actual != expected:
            raise RuntimeError(
                f"Summary formula mismatch in {worksheet.title!r}, row {row}."
            )


def _validate_workbook(
    path: Path,
    template_path: Path,
    source_identities: dict[str, list[tuple[Any, str]]],
) -> dict[str, Any]:
    gateway = WorkbookGateway()
    gateway.verify_openable(path)
    with zipfile.ZipFile(path) as archive:
        bad_member = archive.testzip()
        if bad_member is not None:
            raise RuntimeError(f"Corrupt XLSX member: {bad_member}")

    workbook = load_workbook(path, read_only=False, data_only=False, keep_links=True)
    template = load_workbook(
        template_path,
        read_only=False,
        data_only=False,
        keep_links=True,
    )
    months = MonthSheetService()
    headers = HeaderResolver()
    posting = ExpensePostingService(provider=object())
    details: dict[str, Any] = {}
    try:
        for worksheet in workbook.worksheets:
            if months.parse_target_sheet(worksheet.title) is None:
                continue
            if worksheet.title not in template.sheetnames:
                raise RuntimeError(
                    f"Template is missing month sheet {worksheet.title!r}."
                )
            target = template[worksheet.title]
            target_summary = payment_sync.find_summary_start(target)
            summary_start = payment_sync.find_summary_start(worksheet)
            if summary_start != target_summary:
                raise RuntimeError(
                    f"Summary position differs from template in {worksheet.title!r}."
                )
            status_column = summary_start + len(RPA_SUMMARY_HEADERS) + 1
            actual_headers = tuple(
                worksheet.cell(1, column).value
                for column in range(1, status_column + 1)
            )
            expected_headers = tuple(
                target.cell(1, column).value
                for column in range(1, status_column + 1)
            )
            if actual_headers != expected_headers:
                raise RuntimeError(
                    f"Canonical headers differ in {worksheet.title!r}."
                )

            resolution = headers.resolve(
                worksheet,
                SOURCE_HEADER_ALIASES,
                required=SYNC_FIELDS,
            )
            if set(resolution.columns) != set(SYNC_FIELDS):
                raise RuntimeError(
                    f"Daily-sync columns differ in {worksheet.title!r}: "
                    f"{resolution.columns!r}."
                )

            rows = _record_rows(worksheet)
            identities = [_record_identity(worksheet, row) for row in rows]
            if identities != source_identities[worksheet.title]:
                raise RuntimeError(
                    f"Business-row identities changed in {worksheet.title!r}."
                )
            _validate_formula_block(worksheet, rows)

            base = posting._resolve_base_headers(worksheet)
            fee_columns = posting._resolve_fee_columns(worksheet, base)
            if set(fee_columns) != set(FEE_HEADER_ALIASES):
                missing = sorted(set(FEE_HEADER_ALIASES).difference(fee_columns))
                raise RuntimeError(
                    f"Unresolved fee columns in {worksheet.title!r}: {missing!r}."
                )
            RpaExpenseService._summary_columns(worksheet)
            if posting._update_column(base) != summary_start + 10:
                raise RuntimeError(
                    f"Update column is misplaced in {worksheet.title!r}."
                )
            if RpaExpenseService._optional_status_column(worksheet) != status_column:
                raise RuntimeError(
                    f"Status column is misplaced in {worksheet.title!r}."
                )
            details[worksheet.title] = {
                "records": len(rows),
                "summary_start": get_column_letter(summary_start),
                "status_column": get_column_letter(status_column),
                "fee_columns": len(fee_columns),
                "formula_cells": len(rows) * 10,
            }

        second_pass = payment_sync.normalize_bk_workbook(workbook)
        if second_pass.changed or second_pass.issues:
            raise RuntimeError("BK normalization is not idempotent.")
        for worksheet in workbook.worksheets:
            if months.parse_target_sheet(worksheet.title) is None:
                continue
            if ensure_bang_ke_fee_columns(worksheet, header_row=1):
                raise RuntimeError(
                    f"Fee-column normalization is not idempotent in {worksheet.title!r}."
                )
            if payment_sync.refresh_bk_summary_formulas(worksheet):
                raise RuntimeError(
                    f"Summary normalization is not idempotent in {worksheet.title!r}."
                )

        detail = workbook["_CT_VAN_TAI"]
        if detail.sheet_state != "hidden":
            raise RuntimeError("Carrier detail sheet is not hidden.")
        actual_detail_headers = tuple(
            detail.cell(1, column).value
            for column in range(1, len(DETAIL_HEADERS) + 1)
        )
        if actual_detail_headers != DETAIL_HEADERS:
            raise RuntimeError("Carrier detail headers differ from the canonical schema.")
        if detail.max_row != 1:
            raise RuntimeError("Source-only migration unexpectedly carried system history.")
    finally:
        workbook.close()
        template.close()
    return details


def standardize(source: Path, template: Path, output: Path) -> dict[str, Any]:
    source = source.resolve()
    template = template.resolve()
    output = output.resolve()
    for label, path in (("source", source), ("template", template)):
        if not path.is_file():
            raise FileNotFoundError(f"Missing {label} workbook: {path}")
    if output.exists():
        raise FileExistsError(f"Refusing to overwrite existing output: {output}")
    output.parent.mkdir(parents=True, exist_ok=True)

    source_hash = _sha256(source)
    template_hash = _sha256(template)
    workbook = load_workbook(
        source,
        read_only=False,
        data_only=False,
        keep_links=True,
    )
    template_workbook = load_workbook(
        template,
        read_only=False,
        data_only=False,
        keep_links=True,
    )
    months = MonthSheetService()
    source_identities: dict[str, list[tuple[Any, str]]] = {}
    trimmed_tail_cells: dict[str, int] = {}
    for worksheet in workbook.worksheets:
        if months.parse_target_sheet(worksheet.title) is None:
            continue
        source_identities[worksheet.title] = [
            _record_identity(worksheet, row) for row in _record_rows(worksheet)
        ]
        trimmed_tail_cells[worksheet.title] = _trim_styled_tail(worksheet)

    original_data_rows = payment_sync._data_rows
    payment_sync._data_rows = _fast_data_rows
    temporary_path: Path | None = None
    try:
        report = payment_sync.normalize_bk_workbook(workbook)
        sheet_report: dict[str, Any] = {}
        for worksheet in workbook.worksheets:
            if months.parse_target_sheet(worksheet.title) is None:
                continue
            if worksheet.title not in template_workbook.sheetnames:
                raise RuntimeError(
                    f"Template is missing month sheet {worksheet.title!r}."
                )
            sheet_report[worksheet.title] = _canonicalize_sheet(
                worksheet,
                template_workbook[worksheet.title],
            )

        detail = ensure_bk_detail_sheet(workbook)
        detail.auto_filter.ref = "A1:O1"
        if "_CT_VAN_TAI" in template_workbook.sheetnames:
            template_detail = template_workbook["_CT_VAN_TAI"]
            _copy_header_format(template_detail, detail, len(DETAIL_HEADERS))
            detail.freeze_panes = template_detail.freeze_panes
        detail.sheet_state = "hidden"

        calculation = getattr(workbook, "calculation", None)
        if calculation is not None:
            calculation.fullCalcOnLoad = True
            calculation.forceFullCalc = True
            calculation.calcMode = "auto"

        handle, raw_temp = tempfile.mkstemp(
            prefix=f".{output.stem}_",
            suffix=output.suffix,
            dir=output.parent,
        )
        os.close(handle)
        temporary_path = Path(raw_temp)
        workbook.save(temporary_path)
        workbook.close()
        template_workbook.close()

        validation = _validate_workbook(
            temporary_path,
            template,
            source_identities,
        )
        if _sha256(source) != source_hash or _sha256(template) != template_hash:
            raise RuntimeError("A source/template workbook changed during migration.")
        os.replace(temporary_path, output)
        temporary_path = None
        return {
            "output": str(output),
            "bytes": output.stat().st_size,
            "sha256": _sha256(output),
            "source_sha256_unchanged": source_hash,
            "template_sha256_unchanged": template_hash,
            "total_records": sum(len(rows) for rows in source_identities.values()),
            "trimmed_empty_style_cells": trimmed_tail_cells,
            "normalization_warnings": [
                {
                    "sheet": issue.sheet_name,
                    "row": issue.row,
                    "message": issue.message,
                }
                for issue in report.issues
            ],
            "sheets": sheet_report,
            "validation": validation,
            "system_history_rows": 0,
        }
    finally:
        payment_sync._data_rows = original_data_rows
        try:
            workbook.close()
        except Exception:
            pass
        try:
            template_workbook.close()
        except Exception:
            pass
        if temporary_path is not None:
            temporary_path.unlink(missing_ok=True)


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", required=True, type=Path)
    parser.add_argument("--template", required=True, type=Path)
    parser.add_argument("--output", required=True, type=Path)
    args = parser.parse_args()
    result = standardize(args.source, args.template, args.output)
    print(json.dumps(result, ensure_ascii=False, indent=2, default=str))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
