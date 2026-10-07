"""Read NAM fees from BK and export one carrier's detail to a standalone workbook."""

from __future__ import annotations

import os
import re
import tempfile
import hashlib
import json
from collections import Counter
from dataclasses import dataclass, field, replace
from decimal import Decimal
from pathlib import Path
from typing import Any, Callable, Mapping

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .models import ExcelOperation, PaymentSyncItem, SourceSheetCandidate, WorkbookFingerprint
from .payment_sync import (
    FIELD_LABELS, NAM_FIELDS, NormalizationReport, _ensure_bk_structure, _has_money,
    _source_target_items,
)
from .resolvers import MonthSheetService
from .workbook import WorkbookChangedError, WorkbookGateway, ensure_supported_workbook, workbook_fingerprint


class CarrierExportError(RuntimeError):
    pass


_MONTH_SUFFIX = re.compile(r"(?:\s+BK\s*T?\s*|\s+T\s*)(?:0?[1-9]|1[0-2])$", re.IGNORECASE)
_BAD_FILENAME = re.compile(r'[<>:"/\\|?*\x00-\x1f]')
_MULTIPLE_CARRIERS = re.compile(r"\s*(?:/|;|,|\+|\n)\s*")
_FIELDS = tuple(NAM_FIELDS)
_AMOUNT_COLUMNS = (4, 6, 8, 10, 12, 14, 15, 17, 19)
_HEADERS = (
    "Tháng BK", "QT", "SỐ CONT", "Nâng vỏ", "Số HĐ Nâng vỏ",
    "Nâng hàng", "Số HĐ Nâng hàng", "Hạ vỏ", "Số HĐ Hạ vỏ",
    "VS + D/O", "Số HĐ VS + D/O", "Seal", "Số HĐ Seal", "Làm lệnh",
    "Lưu cont", "Số HĐ Lưu cont", "Sửa chữa", "Số HĐ Sửa chữa",
    "Quá tải", "Số HĐ Quá tải", "N.HA VS D/O LỆNH",
    "Lưu cont + Quá tải", "Sửa chữa Cont", "Tổng tiền",
)


def normalize_carriers(value: Any) -> tuple[str, ...]:
    """Parse one BK cell without guessing that different names are aliases."""
    raw = str(value or "").strip()
    if not raw:
        return ()
    names: dict[str, str] = {}
    for part in _MULTIPLE_CARRIERS.split(raw):
        compact = " ".join(part.split())
        if not compact or compact == "0":
            return ()
        name = _MONTH_SUFFIX.sub("", compact).strip(" ._-")
        if not name:
            return ()
        names.setdefault(name.casefold(), name)
    return tuple(names.values())


def normalize_carrier(value: Any) -> str | None:
    names = normalize_carriers(value)
    return names[0] if len(names) == 1 else None


def _number(value: Any) -> Decimal:
    if value in (None, ""):
        return Decimal(0)
    number = Decimal(str(value))
    if not number.is_finite():
        raise ValueError("Số tiền không hữu hạn.")
    return number


def _entered_amount(value: Any) -> Decimal:
    if isinstance(value, str):
        text = value.strip().replace(" ", "")
        if re.fullmatch(r"\d{1,3}(?:[.,]\d{3})+", text):
            return _number(text.replace(".", "").replace(",", ""))
        return _number(text.replace(",", "."))
    return _number(value)


def _source_signature(sheet_name: str, item: PaymentSyncItem) -> str:
    source = {
        "sheet": sheet_name,
        "rows": item.source_rows,
        "sqt": item.sqt,
        "container": item.container,
        "carrier": item.carrier_value,
        "amounts": {field: str(item.values.get(field)) for field in _FIELDS},
        "invoices": {field: item.invoice_values.get(field) for field in _FIELDS},
    }
    return hashlib.sha256(
        json.dumps(source, ensure_ascii=False, sort_keys=True).encode("utf-8")
    ).hexdigest()


def _row_amounts(row: CarrierExportRow, carrier_key: str) -> list[Decimal]:
    if row.needs_allocation:
        return [
            sum(
                (share.amount for share in row.fee_allocations.get(field, ())
                 if share.carrier_key == carrier_key),
                Decimal(0),
            )
            for field in _FIELDS
        ]
    if row.carrier_key == carrier_key:
        return [_number(row.item.values.get(field)) for field in _FIELDS]
    return [Decimal(0) for _ in _FIELDS]


def _row_invoice(row: CarrierExportRow, field: str, carrier_key: str) -> str | None:
    if row.needs_allocation:
        return next(
            (share.invoice for share in row.fee_allocations.get(field, ())
             if share.carrier_key == carrier_key),
            None,
        )
    return row.item.invoice_values.get(field)


@dataclass(frozen=True, slots=True)
class CarrierFeeShare:
    carrier_key: str
    amount: Decimal
    invoice: str | None = None


@dataclass(frozen=True, slots=True)
class CarrierExportRow:
    sheet_name: str
    month: int
    year: int
    carrier_key: str | None
    item: PaymentSyncItem
    error: str | None = None
    carrier_names: tuple[str, ...] = ()
    fee_allocations: Mapping[str, tuple[CarrierFeeShare, ...]] = field(default_factory=dict)
    source_signature: str = ""

    @property
    def needs_allocation(self) -> bool:
        return len(self.carrier_names) > 1 and self.error is None

    @property
    def allocation_complete(self) -> bool:
        return self.needs_allocation and all(
            field in self.fee_allocations
            for field in _FIELDS
            if _has_money(self.item.values.get(field))
        )


@dataclass(frozen=True, slots=True)
class CarrierCandidate:
    key: str
    name: str
    counts: tuple[int, ...]

    @property
    def total(self) -> int:
        return sum(self.counts)


@dataclass(frozen=True, slots=True)
class CarrierExportPreview:
    source_path: Path
    fingerprint: WorkbookFingerprint
    sheet_names: tuple[str, ...]
    months: tuple[tuple[int, int], ...]
    rows: tuple[CarrierExportRow, ...]
    carriers: tuple[CarrierCandidate, ...]
    unresolved_count: int
    run_id: int | None = None
    allocation_choices: Mapping[str, Mapping[str, list[dict[str, str]]]] = field(default_factory=dict)
    selected_carrier_key: str | None = None

    @property
    def allocation_rows(self) -> tuple[CarrierExportRow, ...]:
        return tuple(row for row in self.rows if row.needs_allocation)

    @property
    def missing_carrier_count(self) -> int:
        return sum(not row.carrier_names and row.error is None for row in self.rows)


@dataclass(frozen=True, slots=True)
class CarrierExportResult:
    operation: ExcelOperation
    target_path: Path
    row_count: int
    total: Decimal
    carrier_name: str
    unresolved_count: int
    run_id: int | None = None

    @property
    def message(self) -> str:
        return f"{self.carrier_name}: {self.row_count} dòng, tổng {self.total:,.0f} đ → {self.target_path}"


class CarrierExportService:
    def __init__(
        self,
        settings: Any | None = None,
        *,
        bk_path: str | Path | None = None,
        gateway: WorkbookGateway | None = None,
        run_repository: Any | None = None,
    ) -> None:
        self.bk_path = Path(bk_path or getattr(settings, "bk_workbook_path", ""))
        self.gateway = gateway or WorkbookGateway()
        self.months = MonthSheetService()
        self.run_repository = run_repository

    def source_sheet_candidates(self) -> list[SourceSheetCandidate]:
        path = ensure_supported_workbook(self.bk_path)
        if not path.is_file():
            raise CarrierExportError(f"Không tìm thấy file BK: {path}")
        book = self.gateway.load(path, read_only=True)
        try:
            return [SourceSheetCandidate(month=parsed[0], source_sheet=name)
                    for name in book.sheetnames
                    if (parsed := self.months.parse_target_sheet(name))]
        finally:
            book.close()

    def analyze(
        self,
        sheet_names: list[str] | tuple[str, ...],
        *,
        resume_run_id: int | None = None,
        progress_callback: Callable[[str], None] | None = None,
    ) -> CarrierExportPreview:
        path = ensure_supported_workbook(self.bk_path)
        if not path.is_file():
            raise CarrierExportError(f"Không tìm thấy file BK: {path}")
        names = tuple(dict.fromkeys(sheet_names))
        if not names:
            raise CarrierExportError("Hãy chọn ít nhất một sheet tháng BK.")
        fingerprint = self.gateway.fingerprint(path)
        book = self.gateway.load(path)
        try:
            selected = []
            for name in names:
                parsed = self.months.parse_target_sheet(name)
                if parsed is None or name not in book:
                    raise CarrierExportError(f"Sheet tháng BK không hợp lệ: {name}")
                selected.append((name, *parsed))
            selected.sort(key=lambda entry: (entry[2], entry[1], book.sheetnames.index(entry[0])))
            normalization = NormalizationReport()
            for name, _, _ in selected:
                _ensure_bk_structure(book[name], normalization)
            rows: list[CarrierExportRow] = []
            for name, month, year in selected:
                if progress_callback:
                    progress_callback(f"Đang đọc {name}…")
                items, errors = _source_target_items(
                    book[name], normalization_issues=normalization.issues
                )
                for item in items["NAM"]:
                    if not any(_has_money(item.values.get(field)) for field in _FIELDS):
                        continue
                    carriers = normalize_carriers(item.carrier_value)
                    key = carriers[0].casefold() if len(carriers) == 1 else None
                    error = errors.get(item.item_id)
                    try:
                        for field in _FIELDS:
                            _number(item.values.get(field))
                    except (ValueError, ArithmeticError):
                        error = error or "Khoản phí có số tiền không hợp lệ."
                    if any(len(values) > 1 for values in item.invoice_candidates.values()):
                        error = error or "Một khoản phí có nhiều số hóa đơn."
                    if error:
                        key = None
                    rows.append(CarrierExportRow(
                        name, month, year, key, item, error, carriers,
                        source_signature=_source_signature(name, item),
                    ))
            self.gateway.assert_unchanged(path, fingerprint, label="File BK")
        finally:
            book.close()
        ordered_names = tuple(entry[0] for entry in selected)
        preview = CarrierExportPreview(
            path, fingerprint, ordered_names,
            tuple((entry[1], entry[2]) for entry in selected),
            tuple(rows), (), 0,
        )
        preview = self._recount(preview)
        if self.run_repository is None:
            return preview
        if resume_run_id is None:
            run = self.run_repository.create_run(
                operation=ExcelOperation.CARRIER_EXPORT,
                status="WAITING_USER",
                source_path=path,
                source_fingerprint=fingerprint,
                sheet_name=", ".join(ordered_names),
            )
            preview = replace(preview, run_id=run.id)
            self._save_preview(preview)
            return preview
        run = self.run_repository.require_by_id(resume_run_id)
        payload = run.item_outcomes if isinstance(run.item_outcomes, dict) else {}
        if (
            run.operation != ExcelOperation.CARRIER_EXPORT.value
            or payload.get("kind") != "carrier_allocation_v1"
            or run.source_path is None
            or run.source_path.resolve() != path.resolve()
            or tuple(payload.get("sheet_names", ())) != ordered_names
        ):
            raise CarrierExportError("Lượt xuất đã chọn không khớp file BK hoặc các sheet đã lưu.")
        preview = replace(
            preview, run_id=run.id,
            selected_carrier_key=payload.get("selected_carrier"),
        )
        saved = payload.get("allocations", {})
        if isinstance(saved, dict):
            valid_choices = {
                row.item.item_id: entry["fields"]
                for row in preview.allocation_rows
                if isinstance((entry := saved.get(row.item.item_id)), dict)
                and entry.get("signature") == row.source_signature
                and isinstance(entry.get("fields"), dict)
            }
            preview = self.apply_allocations(preview, valid_choices, allow_partial=True, persist=False)
        self._save_preview(preview)
        return preview

    @staticmethod
    def _recount(preview: CarrierExportPreview) -> CarrierExportPreview:
        names_by_key: dict[str, str] = {}
        counts: dict[str, Counter[str]] = {}
        unresolved = 0
        for row in preview.rows:
            if row.error or not row.carrier_names:
                unresolved += 1
                continue
            if row.needs_allocation and not row.allocation_complete:
                unresolved += 1
                continue
            keys = (
                {share.carrier_key for shares in row.fee_allocations.values() for share in shares}
                if row.needs_allocation else {row.carrier_key}
            )
            for key in keys:
                if not key or not any(_row_amounts(row, key)):
                    continue
                name = next(name for name in row.carrier_names if name.casefold() == key)
                names_by_key.setdefault(key, name)
                counts.setdefault(key, Counter())[row.sheet_name] += 1
        candidates = tuple(
            CarrierCandidate(key, names_by_key[key], tuple(counts[key][name] for name in preview.sheet_names))
            for key in sorted(names_by_key, key=lambda k: names_by_key[k].casefold())
        )
        return replace(preview, carriers=candidates, unresolved_count=unresolved)

    def _save_preview(self, preview: CarrierExportPreview) -> None:
        if preview.run_id is None or self.run_repository is None:
            return
        allocations = {
            row.item.item_id: {
                "signature": row.source_signature,
                "fields": preview.allocation_choices[row.item.item_id],
            }
            for row in preview.allocation_rows
            if row.item.item_id in preview.allocation_choices
        }
        old = self.run_repository.require_by_id(preview.run_id)
        old_payload = old.item_outcomes if isinstance(old.item_outcomes, dict) else {}
        self.run_repository.update_run(
            preview.run_id,
            source_fingerprint=preview.fingerprint,
            sheet_name=", ".join(preview.sheet_names),
            item_outcomes={
                "kind": "carrier_allocation_v1",
                "sheet_names": list(preview.sheet_names),
                "allocations": allocations,
                "selected_carrier": old_payload.get("selected_carrier"),
            },
        )

    def saved_runs(self, *, limit: int = 20) -> list[Any]:
        if self.run_repository is None:
            return []
        path = self.bk_path.resolve()
        return [
            run for run in self.run_repository.list_runs(
                operation=ExcelOperation.CARRIER_EXPORT,
            )
            if run.source_path is not None
            and run.source_path.resolve() == path
            and isinstance(run.item_outcomes, dict)
            and run.item_outcomes.get("kind") == "carrier_allocation_v1"
        ][:limit]

    def apply_allocations(
        self,
        preview: CarrierExportPreview,
        choices: Mapping[str, Mapping[str, list[dict[str, str]]]],
        *,
        allow_partial: bool = False,
        persist: bool = True,
    ) -> CarrierExportPreview:
        valid_ids = {row.item.item_id for row in preview.allocation_rows}
        if set(choices) - valid_ids:
            raise CarrierExportError("Phân bổ có dòng không thuộc lượt xuất này.")
        updated_rows: list[CarrierExportRow] = []
        accepted: dict[str, dict[str, list[dict[str, str]]]] = {}
        for row in preview.rows:
            if not row.needs_allocation:
                updated_rows.append(row)
                continue
            raw_fields = choices.get(row.item.item_id, {})
            allocations: dict[str, tuple[CarrierFeeShare, ...]] = {}
            accepted_fields: dict[str, list[dict[str, str]]] = {}
            for fee in _FIELDS:
                if not _has_money(row.item.values.get(fee)):
                    continue
                selected = raw_fields.get(fee)
                if not selected:
                    if not allow_partial:
                        raise CarrierExportError(
                            f"QT {row.item.sqt}: chưa phân bổ khoản {FIELD_LABELS.get(fee, fee)}."
                        )
                    continue
                try:
                    shares = self._validate_shares(row, fee, selected)
                except (CarrierExportError, ValueError, ArithmeticError):
                    if not allow_partial:
                        raise
                    continue
                allocations[fee] = shares
                accepted_fields[fee] = [
                    {"carrier_key": share.carrier_key, "amount": str(share.amount),
                     "invoice": share.invoice or ""}
                    for share in shares
                ]
            updated_rows.append(replace(row, fee_allocations=allocations))
            if accepted_fields:
                accepted[row.item.item_id] = accepted_fields
        result = self._recount(replace(
            preview, rows=tuple(updated_rows), allocation_choices=accepted,
        ))
        if persist:
            self._save_preview(result)
        return result

    @staticmethod
    def _validate_shares(
        row: CarrierExportRow, fee: str, raw: Any,
    ) -> tuple[CarrierFeeShare, ...]:
        label = FIELD_LABELS.get(fee, fee)
        if not isinstance(raw, (list, tuple)) or not raw:
            raise CarrierExportError(f"QT {row.item.sqt}: phân bổ {label} không hợp lệ.")
        allowed = {name.casefold() for name in row.carrier_names}
        source_amount = _number(row.item.values.get(fee))
        source_invoice = row.item.invoice_values.get(fee)
        shares: list[CarrierFeeShare] = []
        seen: set[str] = set()
        for entry in raw:
            if not isinstance(entry, Mapping):
                raise CarrierExportError(f"QT {row.item.sqt}: phân bổ {label} không hợp lệ.")
            key = str(entry.get("carrier_key", "")).casefold()
            try:
                amount = _entered_amount(entry.get("amount"))
            except (ValueError, ArithmeticError) as exc:
                raise CarrierExportError(
                    f"QT {row.item.sqt}: số tiền {label} không hợp lệ."
                ) from exc
            if key not in allowed or key in seen or amount <= 0:
                raise CarrierExportError(f"QT {row.item.sqt}: bên VT hoặc số tiền {label} không hợp lệ.")
            seen.add(key)
            invoice = str(entry.get("invoice") or "").strip() or None
            shares.append(CarrierFeeShare(key, amount, invoice))
        if sum((share.amount for share in shares), Decimal(0)) != source_amount:
            raise CarrierExportError(f"QT {row.item.sqt}: tổng phân bổ {label} khác số tiền BK.")
        if len(shares) == 1:
            shares[0] = replace(shares[0], invoice=source_invoice)
        elif source_invoice and any(not share.invoice for share in shares):
            raise CarrierExportError(
                f"QT {row.item.sqt}: hãy xác nhận hóa đơn cho từng phần của {label}."
            )
        return tuple(shares)

    @staticmethod
    def output_filename(preview: CarrierExportPreview, carrier_key: str) -> str:
        carrier = next((c for c in preview.carriers if c.key == carrier_key), None)
        if carrier is None:
            raise CarrierExportError("Bên vận tải đã chọn không có trong các sheet BK.")
        safe = _BAD_FILENAME.sub("_", carrier.name).rstrip(" .")
        if not safe:
            safe = "Ben_VT"
        cross_year = len({year for _, year in preview.months}) > 1
        suffix = "_".join(
            f"T{month:02d}_{year}" if cross_year else f"T{month:02d}"
            for month, year in preview.months
        )
        return f"{safe}_{suffix}.xlsx"

    def export(
        self,
        preview: CarrierExportPreview,
        carrier_key: str,
        output_dir: str | Path,
        *,
        overwrite: bool = False,
        progress_callback: Callable[[str], None] | None = None,
    ) -> CarrierExportResult:
        self.gateway.assert_unchanged(preview.source_path, preview.fingerprint, label="File BK")
        if any(row.needs_allocation and not row.allocation_complete for row in preview.rows):
            raise CarrierExportError("Còn dòng có nhiều bên VT chưa phân bổ đủ khoản phí.")
        carrier = next((c for c in preview.carriers if c.key == carrier_key), None)
        if carrier is None:
            raise CarrierExportError("Bên vận tải đã chọn không còn trong danh sách.")
        rows = [
            row for row in preview.rows
            if not row.error and any(_row_amounts(row, carrier_key))
        ]
        if not rows:
            raise CarrierExportError("Bên vận tải đã chọn không có khoản phí trong Nam.")
        destination = Path(output_dir) / self.output_filename(preview, carrier_key)
        if destination.resolve() == preview.source_path.resolve():
            raise CarrierExportError("Không thể ghi đè file BK.")
        prior_fingerprint = workbook_fingerprint(destination) if destination.exists() else None
        if prior_fingerprint is not None and not overwrite:
            raise FileExistsError(f"File đã tồn tại: {destination}")
        destination.parent.mkdir(parents=True, exist_ok=True)
        run = None
        if self.run_repository is not None:
            if preview.run_id is not None:
                run = self.run_repository.update_run(
                    preview.run_id, status="APPLYING", target_path=destination,
                    target_fingerprint_before=prior_fingerprint,
                )
            else:
                run = self.run_repository.create_run(
                    operation=ExcelOperation.CARRIER_EXPORT,
                    status="APPLYING",
                    source_path=preview.source_path,
                    target_path=destination,
                    source_fingerprint=preview.fingerprint,
                    target_fingerprint_before=prior_fingerprint,
                    sheet_name=", ".join(preview.sheet_names),
                )
        temp_path: Path | None = None
        try:
            if progress_callback:
                progress_callback(f"Đang tạo bảng kê cho {carrier.name}…")
            book = Workbook()
            sheet = book.active
            sheet.title = "Chi tiết phí trong Nam"
            sheet.append(_HEADERS)
            total = Decimal(0)
            exported_amounts: list[list[Decimal]] = []
            for source in rows:
                item = source.item
                amounts = _row_amounts(source, carrier_key)
                exported_amounts.append(amounts)
                amount_total = sum(amounts, Decimal(0))
                total += amount_total
                values: list[Any] = [source.sheet_name, item.sqt, item.container]
                for field, amount in zip(_FIELDS, amounts):
                    values.append(float(amount) if amount else None)
                    if field != "command_fee":
                        values.append(
                            _row_invoice(source, field, carrier_key) if amount else None
                        )
                values += [
                    float(sum(amounts[i] for i in (1, 2, 3, 5))),
                    float(amounts[6] + amounts[8]),
                    float(amounts[7]),
                    float(amount_total),
                ]
                sheet.append(values)
                for index in (5, 7, 9, 11, 13, 16, 18, 20):
                    cell = sheet.cell(sheet.max_row, index)
                    if cell.value is not None:
                        cell.value = str(cell.value)
                        cell.number_format = "@"
            sheet.append(["TỔNG CỘNG", None, None] + [None] * 17 + [
                float(sum((values[i] for values in exported_amounts for i in (1, 2, 3, 5)), Decimal(0))),
                float(sum((values[i] for values in exported_amounts for i in (6, 8)), Decimal(0))),
                float(sum((values[7] for values in exported_amounts), Decimal(0))),
                float(total),
            ])
            self._style_sheet(sheet)
            handle, temp_name = tempfile.mkstemp(prefix=".carrier_export_", suffix=".xlsx", dir=destination.parent)
            os.close(handle)
            temp_path = Path(temp_name)
            book.save(temp_path)
            book.close()
            self._verify(temp_path, len(rows), total)
            self.gateway.assert_unchanged(preview.source_path, preview.fingerprint, label="File BK")
            if prior_fingerprint is None:
                if destination.exists():
                    raise FileExistsError(f"File đích vừa xuất hiện: {destination}")
            elif workbook_fingerprint(destination) != prior_fingerprint:
                raise WorkbookChangedError(label="File đích", path=destination)
            os.replace(temp_path, destination)
            temp_path = None
            result = CarrierExportResult(
                ExcelOperation.CARRIER_EXPORT, destination, len(rows), total,
                carrier.name, preview.unresolved_count, preview.run_id,
            )
            if run is not None:
                payload = run.item_outcomes if isinstance(run.item_outcomes, dict) else {}
                self.run_repository.finish_run(
                    run.id, status="SUCCEEDED", target_fingerprint_after=workbook_fingerprint(destination),
                    total_items=len(preview.rows), changed_items=len(rows), skipped_items=preview.unresolved_count,
                    item_outcomes={
                        **payload, "selected_carrier": carrier.key,
                        "carrier": carrier.name, "total": str(total),
                    },
                )
            return result
        except Exception as exc:
            if run is not None:
                self.run_repository.finish_run(run.id, status="FAILED", error_message=str(exc))
            raise
        finally:
            if temp_path is not None:
                temp_path.unlink(missing_ok=True)

    @staticmethod
    def _style_sheet(sheet: Any) -> None:
        sheet.freeze_panes = "D2"
        sheet.auto_filter.ref = f"A1:X{sheet.max_row - 1}"
        sheet.print_options.horizontalCentered = True
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.fitToWidth = 1
        sheet.print_title_rows = "1:1"
        sheet.print_area = f"A1:X{sheet.max_row}"
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="205781")
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        sheet.row_dimensions[1].height = 42
        widths = {1: 16, 2: 9, 3: 19, 24: 18}
        for index in range(1, 25):
            sheet.column_dimensions[get_column_letter(index)].width = widths.get(index, 16)
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                if cell.column in (*_AMOUNT_COLUMNS, 21, 22, 23, 24):
                    cell.number_format = '#,##0.##;[Red](#,##0.##)'
        for cell in sheet[sheet.max_row]:
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="DCEAF4")

    @staticmethod
    def _verify(path: Path, expected_rows: int, expected_total: Decimal) -> None:
        book = load_workbook(path, read_only=True, data_only=True)
        try:
            if len(book.sheetnames) != 1:
                raise CarrierExportError("File xuất không có đúng một sheet.")
            sheet = book.active
            if sheet.max_row != expected_rows + 2 or tuple(cell.value for cell in sheet[1]) != _HEADERS:
                raise CarrierExportError("File xuất có sai số dòng hoặc sai cột.")
            calculated = sum((_number(sheet.cell(row, 24).value) for row in range(2, sheet.max_row)), Decimal(0))
            footer = _number(sheet.cell(sheet.max_row, 24).value)
            if abs(calculated - expected_total) > Decimal("0.01") or abs(footer - expected_total) > Decimal("0.01"):
                raise CarrierExportError("Tổng tiền trong file xuất không khớp dữ liệu BK.")
        finally:
            book.close()
