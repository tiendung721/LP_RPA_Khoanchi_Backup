"""Cấu trúc cột riêng của luồng bảng kê trên workbook BK."""

from __future__ import annotations

import copy
from typing import Any

from openpyxl.utils import get_column_letter

from .headers import normalize_header


class BangKeColumnError(RuntimeError):
    pass


_INVOICE_HEADERS = {
    normalize_header(value)
    for value in ("Hóa đơn", "Số hóa đơn", "Số HĐ", "HD", "HĐ", "hoá đon")
}


def ensure_bang_ke_fee_columns(
    worksheet: Any,
    *,
    header_row: int,
    strict: bool = True,
) -> bool:
    """Xóa cột VAT và bảo đảm Gia hạn sau Sửa chữa/HĐ."""

    from .payment_sync import (
        _copy_cell_style,
        _delete_column,
        _insert_column,
        find_summary_start,
        refresh_bk_summary_formulas,
    )

    changed = False

    def matches(label: str) -> list[int]:
        summary_start = find_summary_start(worksheet, header_row=header_row)
        boundary = summary_start or worksheet.max_column + 1
        key = normalize_header(label)
        return [
            column
            for column in range(1, boundary)
            if normalize_header(worksheet.cell(header_row, column).value) == key
        ]

    vat_keys = {
        normalize_header("VAT"),
        normalize_header("THUẾ GTGT"),
    }
    summary_start = find_summary_start(worksheet, header_row=header_row)
    boundary = summary_start or worksheet.max_column + 1
    vat_columns = [
        column
        for column in range(1, boundary)
        if normalize_header(worksheet.cell(header_row, column).value) in vat_keys
    ]
    for column in reversed(vat_columns):
        _delete_column(worksheet, column, replacement=None)
        changed = True

    sc_columns = matches("SỬA CHỮA")
    if not sc_columns and not strict:
        if find_summary_start(worksheet, header_row=header_row) is not None:
            refresh_bk_summary_formulas(worksheet)
        return changed
    if len(sc_columns) != 1:
        raise BangKeColumnError(
            f"Sheet {worksheet.title} không xác định duy nhất cột SỬA CHỮA."
        )
    sc_column = sc_columns[0]
    invoice_column = sc_column + 1
    if normalize_header(worksheet.cell(header_row, invoice_column).value) not in _INVOICE_HEADERS:
        raise BangKeColumnError(
            f"Sheet {worksheet.title} thiếu cột HĐ sửa chữa ngay sau SỬA CHỮA."
        )

    insertion = invoice_column + 1
    current = matches("GIA HẠN")
    if len(current) > 1:
        raise BangKeColumnError(f"Sheet {worksheet.title} có nhiều cột GIA HẠN.")
    if not current:
        _insert_column(worksheet, insertion)
        styled_rows = {
            row
            for row, column in getattr(worksheet, "_cells", {})
            if column == sc_column
        }
        styled_rows.add(header_row)
        for row in styled_rows:
            _copy_cell_style(
                worksheet.cell(row, sc_column),
                worksheet.cell(row, insertion),
            )
        source_letter = get_column_letter(sc_column)
        target_letter = get_column_letter(insertion)
        worksheet.column_dimensions[target_letter].width = copy.copy(
            worksheet.column_dimensions[source_letter].width
        )
        worksheet.cell(header_row, insertion).value = "GIA HẠN"
        changed = True
    elif current[0] != insertion:
        raise BangKeColumnError(
            f"Sheet {worksheet.title} có cột GIA HẠN không nằm ngay sau HĐ sửa chữa."
        )
    summary_start = find_summary_start(worksheet, header_row=header_row)
    if summary_start is not None:
        refresh_bk_summary_formulas(worksheet)
    return changed


__all__ = ["BangKeColumnError", "ensure_bang_ke_fee_columns"]
