"""Shared carrier routing, workbook detail ledger, and payment summary helpers."""

from __future__ import annotations

import re
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import datetime
from decimal import Decimal, InvalidOperation
from typing import Any, Iterable, Mapping, Sequence

from openpyxl.styles import Alignment, Font, PatternFill

from app.constants import DAILY_SYNC_CARRIER_FEE_CODES

from .models import CarrierSummaryResult


BK_DETAIL_SHEET = "_CT_VAN_TAI"
PAYMENT_SUMMARY_SHEET = "TONG_HOP_VAN_TAI"
PAYMENT_DETAIL_SHEET = "_CT_TONG_HOP_VT"
CARRIER_SEPARATOR = " / "
UNMAPPED_CARRIER = "CHƯA XÁC ĐỊNH"
MISSING_INVOICE = "CHƯA CÓ HĐ"

SEA_FEE_CODES = frozenset({"CB"})
ROAD_FEE_CODES = frozenset(DAILY_SYNC_CARRIER_FEE_CODES - SEA_FEE_CODES)
NAM_FEE_CODES = frozenset(
    {"NV", "HH", "NH", "HV", "VSDL", "LL", "LC", "SC", "QT"}
)
DAILY_MANAGED_CARRIER_GROUPS = frozenset({"SEA", "ROAD"})

BK_CARRIER_HEADER_ALIASES: dict[str, tuple[str, ...]] = {
    "SEA": ("VT biển", "Vận tải biển"),
    "ROAD": ("VT bộ", "Vận tải bộ"),
    # Giữ alias HP cho luồng BK → Thanh toán cũ; Khoản chi → BK không còn
    # dùng nhóm này để HĐ phí không ghi vào cột VT bộ của file Hàng ngày.
    "HP": ("VT bộ", "Vận tải bộ"),
    "NAM": ("VT TRONG NAM", "Vận tải trong Nam"),
}

DETAIL_HEADERS = (
    "Batch ID",
    "Batch hash",
    "Dòng JSON",
    "Kỳ",
    "Sheet BK",
    "SQT",
    "Container",
    "Nhóm",
    "Mã phí",
    "Số HĐ",
    "Carrier nguồn",
    "Carrier hiệu lực",
    "Số tiền",
    "Hành động",
    "Ngày cập nhật",
)

SUMMARY_HEADERS = (
    "Kỳ",
    "Bên vận tải",
    "Số hóa đơn",
    "Tổng có HĐ",
    "Chưa có HĐ",
    "Tổng tháng",
    "Cảnh báo",
)

INVOICE_SUMMARY_HEADERS = (
    "Kỳ",
    "Bên vận tải",
    "Số HĐ",
    "HH",
    "NV",
    "NH",
    "HV",
    "VSDL",
    "LL",
    "LC",
    "SC",
    "QT",
    "Tổng hóa đơn",
    "Số khoản",
    "Cảnh báo",
)

FEE_CODES = ("HH", "NV", "NH", "HV", "VSDL", "LL", "LC", "SC", "QT")
LEGACY_DETAIL_START = 16  # P
INVOICE_SUMMARY_START = 9  # I
WARNING_FILL = PatternFill("solid", fgColor="FFF2CC")


def carrier_group_for_fee(fee: str | None) -> str | None:
    normalized = str(fee or "").strip().upper()
    if normalized in SEA_FEE_CODES:
        return "SEA"
    if normalized in ROAD_FEE_CODES:
        return "ROAD"
    if normalized in NAM_FEE_CODES:
        return "NAM"
    return None


def carrier_text(value: Any) -> str | None:
    if value in (None, ""):
        return None
    text = " ".join(str(value).strip().split())
    return text or None


def carrier_key(value: Any) -> str | None:
    text = carrier_text(value)
    return text.casefold() if text is not None else None


def split_carriers(value: Any) -> list[str]:
    text = carrier_text(value)
    if text is None:
        return []
    parts = re.split(r"\s*/\s*", text)
    return unique_carriers(parts)


def unique_carriers(values: Iterable[Any]) -> list[str]:
    result: list[str] = []
    seen: set[str] = set()
    for value in values:
        text = carrier_text(value)
        key = carrier_key(text)
        if text is None or key is None or key in seen:
            continue
        seen.add(key)
        result.append(text)
    return result


def join_carriers(*values: Any) -> str | None:
    parts: list[str] = []
    for value in values:
        if isinstance(value, (list, tuple, set, frozenset)):
            parts.extend(unique_carriers(value))
        else:
            parts.extend(split_carriers(value))
    unique = unique_carriers(parts)
    return CARRIER_SEPARATOR.join(unique) if unique else None


def period_for_sheet(sheet_name: str) -> str:
    match = re.fullmatch(r"\s*T(\d{1,2})\s+(\d{2,4})\s*", sheet_name)
    if match is None:
        return sheet_name.strip()
    month = int(match.group(1))
    year_text = match.group(2)
    year = int(year_text) if len(year_text) == 4 else 2000 + int(year_text)
    return f"T{month:02d}/{year}"


def _style_header(worksheet: Any, row: int, start: int, headers: Sequence[str]) -> None:
    fill = PatternFill("solid", fgColor="1F4E78")
    for offset, header in enumerate(headers):
        cell = worksheet.cell(row, start + offset, header)
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = fill


def ensure_bk_detail_sheet(workbook: Any) -> Any:
    if BK_DETAIL_SHEET in workbook.sheetnames:
        worksheet = workbook[BK_DETAIL_SHEET]
    else:
        worksheet = workbook.create_sheet(BK_DETAIL_SHEET)
    worksheet.sheet_state = "hidden"
    _style_header(worksheet, 1, 1, DETAIL_HEADERS)
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = f"A1:O{max(1, worksheet.max_row)}"
    return worksheet


def upsert_bk_detail_rows(workbook: Any, rows: Sequence[Mapping[str, Any]]) -> int:
    if not rows:
        return 0
    worksheet = ensure_bk_detail_sheet(workbook)
    existing: dict[tuple[str, int], int] = {}
    for row in range(2, worksheet.max_row + 1):
        batch_hash = carrier_text(worksheet.cell(row, 2).value)
        source_index = worksheet.cell(row, 3).value
        if batch_hash is not None and isinstance(source_index, int):
            existing[(batch_hash, source_index)] = row
    changed = 0
    for detail in rows:
        batch_hash = str(detail["batch_hash"])
        source_index = int(detail["source_item_index"])
        target_row = existing.get((batch_hash, source_index), worksheet.max_row + 1)
        if target_row > worksheet.max_row:
            existing[(batch_hash, source_index)] = target_row
        values = (
            detail.get("batch_id"),
            batch_hash,
            source_index,
            detail.get("period"),
            detail.get("sheet_name"),
            detail.get("sqt"),
            detail.get("container"),
            detail.get("carrier_group"),
            detail.get("fee"),
            detail.get("invoice_no"),
            detail.get("carrier_source"),
            detail.get("carrier_effective") or UNMAPPED_CARRIER,
            detail.get("amount"),
            detail.get("carrier_action"),
            detail.get("updated_at"),
        )
        before = tuple(worksheet.cell(target_row, column).value for column in range(1, 16))
        if before != values:
            for column, value in enumerate(values, 1):
                worksheet.cell(target_row, column).value = value
            worksheet.cell(target_row, 13).number_format = "#,##0"
            worksheet.cell(target_row, 15).number_format = "dd/mm/yyyy hh:mm:ss"
            changed += 1
    worksheet.auto_filter.ref = f"A1:O{max(1, worksheet.max_row)}"
    return changed


def read_bk_detail_rows(workbook: Any, sheet_name: str) -> list[dict[str, Any]]:
    if BK_DETAIL_SHEET not in workbook.sheetnames:
        return []
    worksheet = workbook[BK_DETAIL_SHEET]
    result: list[dict[str, Any]] = []
    for row in range(2, worksheet.max_row + 1):
        if carrier_text(worksheet.cell(row, 5).value) != sheet_name:
            continue
        values = [worksheet.cell(row, column).value for column in range(1, 16)]
        if values[1] in (None, "") or not isinstance(values[2], int):
            continue
        result.append(dict(zip(DETAIL_HEADERS, values, strict=True)))
    return result


def _detail_values(detail: Mapping[str, Any]) -> tuple[Any, ...]:
    return tuple(detail.get(header) for header in DETAIL_HEADERS)


def _detail_identity(detail: Mapping[str, Any]) -> tuple[str, int] | None:
    batch_hash = carrier_text(detail.get("Batch hash"))
    source_index = detail.get("Dòng JSON")
    if batch_hash is None or isinstance(source_index, bool):
        return None
    try:
        return batch_hash, int(source_index)
    except (TypeError, ValueError):
        return None


def _details_from_sheet(worksheet: Any, *, start_column: int = 1) -> list[dict[str, Any]]:
    headers = tuple(
        carrier_text(worksheet.cell(1, start_column + offset).value)
        for offset in range(len(DETAIL_HEADERS))
    )
    if headers != tuple(DETAIL_HEADERS):
        return []
    result: list[dict[str, Any]] = []
    for row in range(2, worksheet.max_row + 1):
        values = tuple(
            worksheet.cell(row, start_column + offset).value
            for offset in range(len(DETAIL_HEADERS))
        )
        detail = dict(zip(DETAIL_HEADERS, values, strict=True))
        if _detail_identity(detail) is not None:
            result.append(detail)
    return result


def _detail_sort_key(detail: Mapping[str, Any]) -> tuple[Any, ...]:
    period = carrier_text(detail.get("Kỳ")) or ""
    carrier = carrier_text(detail.get("Carrier hiệu lực")) or UNMAPPED_CARRIER
    invoice = carrier_text(detail.get("Số HĐ")) or MISSING_INVOICE
    identity = _detail_identity(detail) or ("", -1)
    return (
        _period_sort_key(period),
        carrier.casefold(),
        invoice.casefold(),
        identity[0],
        identity[1],
    )


def _read_payment_details(workbook: Any) -> list[dict[str, Any]]:
    if PAYMENT_DETAIL_SHEET not in workbook.sheetnames:
        return []
    return _details_from_sheet(workbook[PAYMENT_DETAIL_SHEET])


def _legacy_payment_details(workbook: Any) -> list[dict[str, Any]]:
    if PAYMENT_SUMMARY_SHEET not in workbook.sheetnames:
        return []
    return _details_from_sheet(
        workbook[PAYMENT_SUMMARY_SHEET], start_column=LEGACY_DETAIL_START
    )


def _write_payment_details(workbook: Any, details: Sequence[Mapping[str, Any]]) -> bool:
    created = PAYMENT_DETAIL_SHEET not in workbook.sheetnames
    worksheet = (
        workbook.create_sheet(PAYMENT_DETAIL_SHEET)
        if created
        else workbook[PAYMENT_DETAIL_SHEET]
    )
    desired = [tuple(DETAIL_HEADERS), *(_detail_values(detail) for detail in details)]
    before = [
        tuple(worksheet.cell(row, column).value for column in range(1, 16))
        for row in range(1, max(worksheet.max_row, 1) + 1)
    ]
    changed = created or worksheet.sheet_state != "hidden" or before != desired
    worksheet.delete_rows(1, max(worksheet.max_row, 1))
    _style_header(worksheet, 1, 1, DETAIL_HEADERS)
    for row, detail in enumerate(details, 2):
        for column, value in enumerate(_detail_values(detail), 1):
            worksheet.cell(row, column).value = value
        worksheet.cell(row, 13).number_format = "#,##0"
        worksheet.cell(row, 15).number_format = "dd/mm/yyyy hh:mm:ss"
    worksheet.sheet_state = "hidden"
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = f"A1:O{max(1, len(details) + 1)}"
    return changed


def _normalized_key(value: Any) -> str:
    return (carrier_text(value) or "").casefold()


def _amount_decimal(value: Any) -> Decimal:
    if isinstance(value, bool) or value in (None, ""):
        return Decimal(0)
    if isinstance(value, str):
        value = value.replace(",", "").strip()
    try:
        return Decimal(str(value))
    except (InvalidOperation, TypeError, ValueError):
        return Decimal(0)


def _excel_number(value: Decimal) -> int | float:
    integral = value.to_integral_value()
    return int(integral) if value == integral else float(value)


def _period_sort_key(period: str) -> tuple[Any, ...]:
    match = re.fullmatch(r"T(\d{1,2})/(\d{4})", period.strip(), re.IGNORECASE)
    if match is None:
        return (1, period.casefold())
    return (0, -int(match.group(2)), -int(match.group(1)))


@dataclass(slots=True)
class _InvoiceAggregate:
    period: str
    carrier: str
    invoice: str
    missing_invoice: bool
    fees: dict[str, Decimal] = field(
        default_factory=lambda: {fee: Decimal(0) for fee in FEE_CODES}
    )
    total: Decimal = Decimal(0)
    item_count: int = 0
    warnings: set[str] = field(default_factory=set)


@dataclass(slots=True)
class _MonthlyAggregate:
    period: str
    carrier: str
    invoice_count: int = 0
    invoiced_total: Decimal = Decimal(0)
    missing_invoice_total: Decimal = Decimal(0)
    warnings: set[str] = field(default_factory=set)

    @property
    def total(self) -> Decimal:
        return self.invoiced_total + self.missing_invoice_total


@dataclass(slots=True)
class _SummaryBuild:
    monthly_rows: list[tuple[Any, ...]]
    invoice_rows: list[tuple[Any, ...]]
    carrier_count: int
    invoice_count: int
    selected_period_total: int | float
    missing_invoice_items: int
    unmapped_carrier_items: int
    cross_carrier_invoices: int
    suspected_duplicate_items: int


def _build_summary(
    details: Sequence[Mapping[str, Any]], *, selected_period: str
) -> _SummaryBuild:
    prepared: list[dict[str, Any]] = []
    missing_invoice_items = 0
    unmapped_carrier_items = 0
    for detail in details:
        period = carrier_text(detail.get("Kỳ")) or period_for_sheet(
            carrier_text(detail.get("Sheet BK")) or ""
        )
        carrier = carrier_text(detail.get("Carrier hiệu lực")) or UNMAPPED_CARRIER
        invoice_value = carrier_text(detail.get("Số HĐ"))
        missing_invoice = invoice_value is None
        invoice = invoice_value or MISSING_INVOICE
        if missing_invoice:
            missing_invoice_items += 1
        if carrier == UNMAPPED_CARRIER:
            unmapped_carrier_items += 1
        prepared.append(
            {
                "period": period,
                "period_key": _normalized_key(period),
                "carrier": carrier,
                "carrier_key": _normalized_key(carrier),
                "invoice": invoice,
                "invoice_key": _normalized_key(invoice),
                "missing_invoice": missing_invoice,
                "fee": str(detail.get("Mã phí") or "").strip().upper(),
                "container_key": _normalized_key(detail.get("Container")),
                "amount": _amount_decimal(detail.get("Số tiền")),
                "batch_hash": carrier_text(detail.get("Batch hash")) or "",
            }
        )

    invoice_carriers: dict[tuple[str, str], set[str]] = defaultdict(set)
    for item in prepared:
        if not item["missing_invoice"]:
            invoice_carriers[(item["period_key"], item["invoice_key"])].add(
                item["carrier_key"]
            )
    cross_invoice_keys = {
        key for key, carriers in invoice_carriers.items() if len(carriers) > 1
    }

    duplicate_groups: dict[tuple[Any, ...], list[int]] = defaultdict(list)
    for index, item in enumerate(prepared):
        duplicate_groups[
            (
                item["period_key"],
                item["carrier_key"],
                item["invoice_key"],
                item["container_key"],
                item["fee"],
                item["amount"],
            )
        ].append(index)
    duplicate_indexes: set[int] = set()
    suspected_duplicate_items = 0
    for indexes in duplicate_groups.values():
        if len({prepared[index]["batch_hash"] for index in indexes}) <= 1:
            continue
        duplicate_indexes.update(indexes)
        suspected_duplicate_items += len(indexes) - 1

    invoices: dict[tuple[str, str, str], _InvoiceAggregate] = {}
    duplicate_counts_by_invoice: dict[tuple[str, str, str], int] = defaultdict(int)
    for index, item in enumerate(prepared):
        key = (item["period_key"], item["carrier_key"], item["invoice_key"])
        aggregate = invoices.setdefault(
            key,
            _InvoiceAggregate(
                period=item["period"],
                carrier=item["carrier"],
                invoice=item["invoice"],
                missing_invoice=item["missing_invoice"],
            ),
        )
        aggregate.total += item["amount"]
        aggregate.item_count += 1
        if item["fee"] in aggregate.fees:
            aggregate.fees[item["fee"]] += item["amount"]
        else:
            aggregate.warnings.add(f"Mã phí chưa hỗ trợ: {item['fee'] or 'trống'}")
        if aggregate.missing_invoice:
            aggregate.warnings.add(MISSING_INVOICE)
        if aggregate.carrier == UNMAPPED_CARRIER:
            aggregate.warnings.add(UNMAPPED_CARRIER)
        if (item["period_key"], item["invoice_key"]) in cross_invoice_keys:
            aggregate.warnings.add("Hóa đơn thuộc nhiều bên vận tải")
        if index in duplicate_indexes:
            duplicate_counts_by_invoice[key] += 1

    for key, count in duplicate_counts_by_invoice.items():
        invoices[key].warnings.add(f"Có {count} khoản trong nhóm nghi trùng")

    monthly: dict[tuple[str, str], _MonthlyAggregate] = {}
    for invoice in invoices.values():
        key = (_normalized_key(invoice.period), _normalized_key(invoice.carrier))
        aggregate = monthly.setdefault(
            key, _MonthlyAggregate(period=invoice.period, carrier=invoice.carrier)
        )
        if invoice.missing_invoice:
            aggregate.missing_invoice_total += invoice.total
        else:
            aggregate.invoice_count += 1
            aggregate.invoiced_total += invoice.total
        aggregate.warnings.update(invoice.warnings)

    monthly_values = sorted(
        monthly.values(), key=lambda value: (_period_sort_key(value.period), value.carrier.casefold())
    )
    invoice_values = sorted(
        invoices.values(),
        key=lambda value: (
            _period_sort_key(value.period),
            value.carrier.casefold(),
            value.missing_invoice,
            value.invoice.casefold(),
        ),
    )
    monthly_rows = [
        (
            value.period,
            value.carrier,
            value.invoice_count,
            _excel_number(value.invoiced_total),
            _excel_number(value.missing_invoice_total),
            _excel_number(value.total),
            "; ".join(sorted(value.warnings)) or None,
        )
        for value in monthly_values
    ]
    invoice_rows = [
        (
            value.period,
            value.carrier,
            value.invoice,
            *(_excel_number(value.fees[fee]) for fee in FEE_CODES),
            _excel_number(value.total),
            value.item_count,
            "; ".join(sorted(value.warnings)) or None,
        )
        for value in invoice_values
    ]
    selected_key = _normalized_key(selected_period)
    selected_monthly = [
        value for value in monthly_values if _normalized_key(value.period) == selected_key
    ]
    return _SummaryBuild(
        monthly_rows=monthly_rows,
        invoice_rows=invoice_rows,
        carrier_count=len(selected_monthly),
        invoice_count=sum(value.invoice_count for value in selected_monthly),
        selected_period_total=_excel_number(
            sum((value.total for value in selected_monthly), Decimal(0))
        ),
        missing_invoice_items=sum(
            1 for item in prepared if item["missing_invoice"] and item["period_key"] == selected_key
        ),
        unmapped_carrier_items=sum(
            1
            for item in prepared
            if item["carrier"] == UNMAPPED_CARRIER and item["period_key"] == selected_key
        ),
        cross_carrier_invoices=sum(1 for period, _invoice in cross_invoice_keys if period == selected_key),
        suspected_duplicate_items=sum(
            len(indexes) - 1
            for indexes in duplicate_groups.values()
            if indexes
            and prepared[indexes[0]]["period_key"] == selected_key
            and len({prepared[index]["batch_hash"] for index in indexes}) > 1
        ),
    )


def _summary_matrix(build: _SummaryBuild) -> list[tuple[Any, ...]]:
    height = max(len(build.monthly_rows), len(build.invoice_rows)) + 1
    width = INVOICE_SUMMARY_START - 1 + len(INVOICE_SUMMARY_HEADERS)
    matrix: list[list[Any]] = [[None] * width for _ in range(height)]
    matrix[0][: len(SUMMARY_HEADERS)] = SUMMARY_HEADERS
    start = INVOICE_SUMMARY_START - 1
    matrix[0][start : start + len(INVOICE_SUMMARY_HEADERS)] = INVOICE_SUMMARY_HEADERS
    for row, values in enumerate(build.monthly_rows, 1):
        matrix[row][: len(values)] = values
    for row, values in enumerate(build.invoice_rows, 1):
        matrix[row][start : start + len(values)] = values
    return [tuple(row) for row in matrix]


def _write_summary_sheet(workbook: Any, build: _SummaryBuild) -> bool:
    created = PAYMENT_SUMMARY_SHEET not in workbook.sheetnames
    worksheet = (
        workbook.create_sheet(PAYMENT_SUMMARY_SHEET)
        if created
        else workbook[PAYMENT_SUMMARY_SHEET]
    )
    desired = _summary_matrix(build)
    compare_rows = max(worksheet.max_row, len(desired), 1)
    compare_columns = max(worksheet.max_column, len(desired[0]), LEGACY_DETAIL_START + len(DETAIL_HEADERS) - 1)
    before = [
        tuple(worksheet.cell(row, column).value for column in range(1, compare_columns + 1))
        for row in range(1, compare_rows + 1)
    ]
    padded_desired = [
        tuple(
            desired[row - 1][column - 1]
            if row <= len(desired) and column <= len(desired[0])
            else None
            for column in range(1, compare_columns + 1)
        )
        for row in range(1, compare_rows + 1)
    ]
    changed = created or worksheet.sheet_state != "visible" or before != padded_desired
    for merged in tuple(worksheet.merged_cells.ranges):
        worksheet.unmerge_cells(str(merged))
    for row in range(1, compare_rows + 1):
        for column in range(1, compare_columns + 1):
            cell = worksheet.cell(row, column)
            cell.value = None
            cell.fill = PatternFill(fill_type=None)
            cell.font = Font()
            cell.alignment = Alignment()
            cell.number_format = "General"
    for row, values in enumerate(desired, 1):
        for column, value in enumerate(values, 1):
            worksheet.cell(row, column).value = value
    _style_header(worksheet, 1, 1, SUMMARY_HEADERS)
    _style_header(worksheet, 1, INVOICE_SUMMARY_START, INVOICE_SUMMARY_HEADERS)
    for row in range(2, len(build.monthly_rows) + 2):
        for column in range(3, 7):
            worksheet.cell(row, column).number_format = "#,##0"
        if worksheet.cell(row, 7).value:
            worksheet.cell(row, 7).fill = WARNING_FILL
    invoice_total_column = INVOICE_SUMMARY_START + 3 + len(FEE_CODES)
    for row in range(2, len(build.invoice_rows) + 2):
        for column in range(INVOICE_SUMMARY_START + 3, invoice_total_column + 1):
            worksheet.cell(row, column).number_format = "#,##0"
        warning_column = INVOICE_SUMMARY_START + len(INVOICE_SUMMARY_HEADERS) - 1
        if worksheet.cell(row, warning_column).value:
            worksheet.cell(row, warning_column).fill = WARNING_FILL
    for row in worksheet.iter_rows(
        min_row=1,
        max_row=max(len(desired), 1),
        min_col=1,
        max_col=len(desired[0]),
    ):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    widths = {
        "A": 12,
        "B": 24,
        "C": 12,
        "D": 16,
        "E": 16,
        "F": 16,
        "G": 40,
        "I": 12,
        "J": 24,
        "K": 18,
        "L": 14,
        "M": 14,
        "N": 14,
        "O": 14,
        "P": 14,
        "Q": 14,
        "R": 14,
        "S": 14,
        "T": 14,
        "U": 18,
        "V": 12,
        "W": 42,
    }
    for column, width in widths.items():
        worksheet.column_dimensions[column].width = width
    worksheet.sheet_state = "visible"
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = f"A1:G{max(1, len(build.monthly_rows) + 1)}"
    return changed


def sync_payment_summary_sheet(
    workbook: Any,
    details: Sequence[Mapping[str, Any]],
    *,
    source_sheet: str,
) -> CarrierSummaryResult:
    """Thay dữ liệu một sheet nguồn rồi dựng lại tổng hợp tích lũy hai tầng."""

    existing = _read_payment_details(workbook)
    legacy = _legacy_payment_details(workbook)
    managed_exists = bool(
        existing
        or legacy
        or PAYMENT_DETAIL_SHEET in workbook.sheetnames
        or PAYMENT_SUMMARY_SHEET in workbook.sheetnames
        or details
    )
    selected_key = _normalized_key(source_sheet)
    period = period_for_sheet(source_sheet)
    selected_period_key = _normalized_key(period)
    merged: dict[tuple[str, int], dict[str, Any]] = {}
    for detail in (*legacy, *existing):
        detail_period = carrier_text(detail.get("Kỳ")) or period_for_sheet(
            carrier_text(detail.get("Sheet BK")) or ""
        )
        if (
            _normalized_key(detail.get("Sheet BK")) == selected_key
            or _normalized_key(detail_period) == selected_period_key
        ):
            continue
        identity = _detail_identity(detail)
        if identity is not None:
            merged[identity] = dict(detail)
    for detail in details:
        identity = _detail_identity(detail)
        if identity is not None:
            merged[identity] = dict(detail)
    combined = sorted(merged.values(), key=_detail_sort_key)
    build = _build_summary(combined, selected_period=period)
    if not managed_exists and not combined:
        return CarrierSummaryResult(source_sheet_name=source_sheet, period=period)
    ledger_changed = _write_payment_details(workbook, combined)
    summary_changed = _write_summary_sheet(workbook, build)
    return CarrierSummaryResult(
        changed=ledger_changed or summary_changed or bool(legacy),
        source_sheet_name=source_sheet,
        period=period,
        carrier_count=build.carrier_count,
        invoice_count=build.invoice_count,
        selected_period_total=build.selected_period_total,
        missing_invoice_items=build.missing_invoice_items,
        unmapped_carrier_items=build.unmapped_carrier_items,
        cross_carrier_invoices=build.cross_carrier_invoices,
        suspected_duplicate_items=build.suspected_duplicate_items,
        detail_count=len(combined),
    )


def detail_rows_from_actions(
    actions: Sequence[Mapping[str, Any]],
    *,
    batch_id: int | None,
    batch_hash: str,
    sheet_name: str,
    updated_at: datetime,
) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    period = period_for_sheet(sheet_name)
    for action in actions:
        if action.get("carrier_group") not in {"HP", "NAM"}:
            continue
        sources = {
            int(source["source_item_index"]): source
            for source in action.get("source_items", ())
        }
        for source_index in action.get("source_indices", ()):
            source = sources.get(int(source_index), {})
            rows.append(
                {
                    "batch_id": batch_id,
                    "batch_hash": batch_hash,
                    "source_item_index": int(source_index),
                    "period": period,
                    "sheet_name": sheet_name,
                    "sqt": action.get("source_sqt"),
                    "container": source.get("container", action.get("container")),
                    "carrier_group": action.get("carrier_group"),
                    "fee": source.get("fee", action.get("fee_selected")),
                    "invoice_no": source.get("invoice_no") or action.get("invoice_selected"),
                    "carrier_source": source.get("carrier"),
                    "carrier_effective": action.get("carrier_effective"),
                    "amount": source.get("amount", action.get("value_after")),
                    "carrier_action": str(
                        getattr(action.get("carrier_action"), "value", action.get("carrier_action") or "")
                    ),
                    "updated_at": updated_at,
                }
            )
    return rows


__all__ = [
    "BK_CARRIER_HEADER_ALIASES",
    "BK_DETAIL_SHEET",
    "DAILY_MANAGED_CARRIER_GROUPS",
    "DETAIL_HEADERS",
    "INVOICE_SUMMARY_HEADERS",
    "MISSING_INVOICE",
    "PAYMENT_DETAIL_SHEET",
    "PAYMENT_SUMMARY_SHEET",
    "SUMMARY_HEADERS",
    "UNMAPPED_CARRIER",
    "carrier_group_for_fee",
    "carrier_key",
    "carrier_text",
    "detail_rows_from_actions",
    "ensure_bk_detail_sheet",
    "join_carriers",
    "read_bk_detail_rows",
    "split_carriers",
    "sync_payment_summary_sheet",
    "unique_carriers",
    "upsert_bk_detail_rows",
]
