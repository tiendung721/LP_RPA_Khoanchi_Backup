"""Export the BK rows changed by the latest completed expense posting run."""

from __future__ import annotations

import os
import re
import tempfile
from collections import defaultdict
from dataclasses import dataclass
from datetime import date, datetime, time
from decimal import Decimal
from pathlib import Path
from typing import Any, Callable, Mapping, Sequence

from openpyxl import Workbook, load_workbook
from openpyxl.formula.translate import Translator
from openpyxl.styles import Alignment, Border, Color, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .headers import normalize_header
from .models import ExcelOperation, PostingPlan, PostingResult
from .payment_sync import find_summary_start
from .workbook import WorkbookChangedError, workbook_fingerprint


class PostingExportError(RuntimeError):
    pass


SNAPSHOT_VERSION = 2
CHANGED_FILL = "FFF2B3"
CHANGED_BORDER = "D59D00"
_TECHNICAL_HEADERS = {normalize_header("Date cập nhật"), normalize_header("Trạng thái RPA")}


def _business_last_column(sheet: Any) -> int:
    """Stop before BK's calculated summary and technical columns."""

    summary_start = find_summary_start(sheet)
    if summary_start is not None:
        return summary_start - 1
    for column in range(1, sheet.max_column + 1):
        if normalize_header(sheet.cell(1, column).value) in _TECHNICAL_HEADERS:
            return column - 1
    return max(
        (column for column in range(1, sheet.max_column + 1) if sheet.cell(1, column).value is not None),
        default=0,
    )


def _encode_cell(cell: Any) -> dict[str, Any] | None:
    value = cell.value
    if value is None:
        return None
    if cell.data_type == "f":
        return {"kind": "formula", "value": str(value), "origin": cell.coordinate}
    if isinstance(value, datetime):
        return {"kind": "datetime", "value": value.isoformat()}
    if isinstance(value, date):
        return {"kind": "date", "value": value.isoformat()}
    if isinstance(value, time):
        return {"kind": "time", "value": value.isoformat()}
    if isinstance(value, Decimal):
        return {"kind": "decimal", "value": str(value)}
    if isinstance(value, (str, int, float, bool)):
        return {"kind": "value", "value": value}
    return {"kind": "value", "value": str(value)}


def _decoded_value(saved: Mapping[str, Any] | None, coordinate: str) -> Any:
    if saved is None:
        return None
    kind = saved["kind"]
    value = saved["value"]
    if kind == "formula":
        return Translator(value, origin=saved["origin"]).translate_formula(coordinate)
    if kind == "datetime":
        return datetime.fromisoformat(value)
    if kind == "date":
        return date.fromisoformat(value)
    if kind == "time":
        return time.fromisoformat(value)
    if kind == "decimal":
        return Decimal(value)
    return value


def _write_cell(cell: Any, saved: Mapping[str, Any] | None) -> None:
    cell.value = _decoded_value(saved, cell.coordinate)
    if saved is not None and saved["kind"] == "value" and isinstance(saved["value"], str):
        # Invoice numbers and other source text must never become formulas.
        cell.data_type = "s"


def _color_spec(color: Any) -> dict[str, Any] | None:
    if color is None or color.type not in {"rgb", "indexed", "theme", "auto"}:
        return None
    return {"type": color.type, "value": getattr(color, color.type), "tint": color.tint}


def _color(saved: Mapping[str, Any] | None) -> Color | None:
    if saved is None:
        return None
    return Color(**{saved["type"]: saved["value"]}, tint=saved["tint"])


def _header_style(cell: Any) -> dict[str, Any]:
    return {
        "font": {
            "name": cell.font.name,
            "size": cell.font.sz,
            "bold": cell.font.bold,
            "italic": cell.font.italic,
            "color": _color_spec(cell.font.color),
        },
        "fill": {
            "pattern": cell.fill.patternType,
            "foreground": _color_spec(cell.fill.fgColor),
            "background": _color_spec(cell.fill.bgColor),
        },
        "alignment": {
            "horizontal": cell.alignment.horizontal,
            "vertical": cell.alignment.vertical,
            "wrap_text": cell.alignment.wrap_text,
        },
    }


def _apply_header_style(cell: Any, saved: Mapping[str, Any]) -> None:
    font = saved["font"]
    cell.font = Font(
        name=font["name"], size=font["size"], bold=font["bold"],
        italic=font["italic"], color=_color(font["color"]),
    )
    fill = saved["fill"]
    cell.fill = PatternFill(
        patternType=fill["pattern"], fgColor=_color(fill["foreground"]),
        bgColor=_color(fill["background"]),
    )
    alignment = saved["alignment"]
    cell.alignment = Alignment(
        horizontal=alignment["horizontal"], vertical=alignment["vertical"],
        wrap_text=alignment["wrap_text"],
    )


def _changed_columns(action: Mapping[str, Any]) -> set[int]:
    columns: set[int] = set()
    for written, target in (
        ("amount_write", "target_column"),
        ("invoice_write", "invoice_target_column"),
        ("carrier_write", "carrier_target_column"),
    ):
        if action.get(written):
            value = action.get(target)
            if value is None or int(value) < 1:
                raise PostingExportError(f"Thiếu cột BK đã ghi: {target}.")
            columns.add(int(value))
    return columns


def build_posting_export_snapshot(
    plan: PostingPlan,
    actions: Sequence[Mapping[str, Any]],
    result: PostingResult,
) -> dict[str, Any]:
    """Capture final BK values now, so a later export cannot drift with the BK file."""

    changed: dict[str, dict[int, set[int]]] = defaultdict(lambda: defaultdict(set))
    for action in actions:
        columns = _changed_columns(action)
        if not columns:
            continue
        sheet_name = str(action.get("sheet_name") or "")
        target_row = action.get("target_row")
        if not sheet_name or target_row is None or int(target_row) < 2:
            raise PostingExportError("Khoản đã ghi thiếu sheet hoặc dòng BK.")
        changed[sheet_name][int(target_row)].update(columns)

    sheets: list[dict[str, Any]] = []
    sqt_keys: set[tuple[str, str]] = set()
    if changed:
        book = load_workbook(result.target_path, read_only=False, data_only=False)
        try:
            missing = set(changed).difference(book.sheetnames)
            if missing:
                raise PostingExportError(f"Không tìm thấy sheet BK đã ghi: {', '.join(sorted(missing))}.")
            for name in book.sheetnames:
                if name not in changed:
                    continue
                sheet = book[name]
                last_column = _business_last_column(sheet)
                if last_column < 1:
                    raise PostingExportError(f"Sheet {name} không có header nghiệp vụ BK.")
                headers = [sheet.cell(1, col).value for col in range(1, last_column + 1)]
                if normalize_header(headers[0]) not in {"sqt", "sqt pm", "so thu tu pm"}:
                    raise PostingExportError(f"Sheet {name} không bắt đầu bằng cột SQT.")
                header_styles = [_header_style(sheet.cell(1, col)) for col in range(1, last_column + 1)]
                widths = [
                    sheet.column_dimensions[get_column_letter(col)].width or 15
                    for col in range(1, last_column + 1)
                ]
                rows: list[dict[str, Any]] = []
                for bk_row, columns in sorted(changed[name].items()):
                    if bk_row > sheet.max_row:
                        raise PostingExportError(f"Dòng {bk_row} không còn trong sheet {name}.")
                    if any(col > last_column for col in columns):
                        raise PostingExportError(f"Ô vừa ghi nằm ngoài cột nghiệp vụ của sheet {name}.")
                    cells = [sheet.cell(bk_row, col) for col in range(1, last_column + 1)]
                    sqt = cells[0].value
                    if sqt is None or str(sqt).strip() == "":
                        raise PostingExportError(f"Dòng {bk_row} trên sheet {name} không có SQT.")
                    sqt_keys.add((name, str(sqt).strip()))
                    rows.append({
                        "bk_row": bk_row,
                        "sqt": str(sqt),
                        "values": [_encode_cell(cell) for cell in cells],
                        "formats": [cell.number_format for cell in cells],
                        "changed_columns": sorted(columns),
                    })
                sheets.append({
                    "name": name, "headers": headers, "header_styles": header_styles,
                    "widths": widths, "rows": rows,
                })
        finally:
            book.close()
    return {
        "version": SNAPSHOT_VERSION,
        "run_id": result.run_id,
        "completed_at": datetime.now().astimezone().isoformat(timespec="seconds"),
        "bk_path": str(result.target_path),
        "source_path": str(plan.batch_path),
        "row_count": sum(len(sheet["rows"]) for sheet in sheets),
        "sqt_count": len(sqt_keys),
        "sheets": sheets,
    }


@dataclass(frozen=True, slots=True)
class PostingExportPreview:
    run_id: int
    completed_at: str
    row_count: int
    sqt_count: int
    sheet_count: int
    filename: str
    last_export_path: Path | None = None


@dataclass(frozen=True, slots=True)
class PostingExportResult:
    target_path: Path
    run_id: int
    row_count: int
    sqt_count: int
    sheet_count: int
    operation: ExcelOperation = ExcelOperation.POSTING_EXPORT

    @property
    def message(self) -> str:
        return f"{self.sqt_count} SQT, {self.row_count} dòng BK → {self.target_path}"


class PostingExportService:
    def __init__(self, run_repository: Any, posting_repository: Any) -> None:
        self.run_repository = run_repository
        self.posting_repository = posting_repository

    def latest_preview(self) -> PostingExportPreview | None:
        run = self.run_repository.get_latest(
            operation=ExcelOperation.EXPENSE_POSTING,
            statuses=("SUCCEEDED", "NO_CHANGES"),
        )
        if run is None:
            return None
        snapshot = self.posting_repository.get_export_snapshot(run.id)
        if snapshot is None or snapshot.get("version") != SNAPSHOT_VERSION:
            return None
        completed = str(snapshot["completed_at"])
        stamp = datetime.fromisoformat(completed.replace("Z", "+00:00")).strftime("%Y%m%d_%H%M")
        names = [sheet["name"] for sheet in snapshot["sheets"]]
        if 1 <= len(names) <= 3:
            period = "_".join(
                re.sub(r"[^A-Za-z0-9]+", "-", name).strip("-")
                for name in names
            )
        else:
            period = f"{len(names)}thang"
        filename = f"BK_vua_nhap_{period}_{snapshot['sqt_count']}SQT_{stamp}_L{run.id}.xlsx"
        return PostingExportPreview(
            run_id=run.id,
            completed_at=completed,
            row_count=int(snapshot["row_count"]),
            sqt_count=int(snapshot["sqt_count"]),
            sheet_count=len(names),
            filename=filename,
            last_export_path=self.posting_repository.get_export_path(run.id),
        )

    def export(
        self,
        run_id: int,
        destination: str | Path,
        *,
        overwrite: bool = False,
        progress_callback: Callable[[str], None] | None = None,
    ) -> PostingExportResult:
        run = self.run_repository.require_by_id(run_id)
        if run.operation != ExcelOperation.EXPENSE_POSTING.value or run.status not in {"SUCCEEDED", "NO_CHANGES"}:
            raise PostingExportError("Lượt nhập BK chưa hoàn tất.")
        snapshot = self.posting_repository.get_export_snapshot(run_id)
        if snapshot is None or snapshot.get("version") != SNAPSHOT_VERSION:
            raise PostingExportError("Lượt nhập này chưa có ảnh chụp dòng BK theo format mới.")
        if int(snapshot["row_count"]) == 0:
            raise PostingExportError("Lượt này không có dòng BK được cập nhật.")
        target = Path(destination)
        if target.suffix.lower() != ".xlsx":
            raise PostingExportError("File xuất phải có đuôi .xlsx.")
        if run.target_path is not None and target.resolve() == Path(run.target_path).resolve():
            raise PostingExportError("Không thể ghi đè file BK.")
        before = workbook_fingerprint(target) if target.exists() else None
        if before is not None and not overwrite:
            raise FileExistsError(f"File đã tồn tại: {target}")
        target.parent.mkdir(parents=True, exist_ok=True)
        temp_path: Path | None = None
        try:
            if progress_callback:
                progress_callback("Đang dựng các dòng BK vừa nhập…")
            book = self._workbook(snapshot)
            handle, name = tempfile.mkstemp(prefix=".posting_export_", suffix=".xlsx", dir=target.parent)
            os.close(handle)
            temp_path = Path(name)
            try:
                book.save(temp_path)
            finally:
                book.close()
            self._verify(temp_path, snapshot)
            if before is None:
                if target.exists():
                    raise FileExistsError(f"File đích vừa xuất hiện: {target}")
            elif workbook_fingerprint(target) != before:
                raise WorkbookChangedError(label="File đích", path=target)
            os.replace(temp_path, target)
            temp_path = None
            self.posting_repository.set_export_path(run_id, target)
            return PostingExportResult(
                target, run_id, int(snapshot["row_count"]),
                int(snapshot["sqt_count"]), len(snapshot["sheets"]),
            )
        finally:
            if temp_path is not None:
                temp_path.unlink(missing_ok=True)

    @staticmethod
    def _workbook(snapshot: Mapping[str, Any]) -> Workbook:
        book = Workbook()
        book.remove(book.active)
        book.calculation.fullCalcOnLoad = True
        book.calculation.forceFullCalc = True
        book.calculation.calcMode = "auto"
        for source in snapshot["sheets"]:
            sheet = book.create_sheet(source["name"])
            end = get_column_letter(len(source["headers"]))
            sheet.sheet_view.showGridLines = False
            for col, (header, width, style) in enumerate(
                zip(source["headers"], source["widths"], source["header_styles"]), 1
            ):
                cell = sheet.cell(1, col, header)
                _apply_header_style(cell, style)
                sheet.column_dimensions[get_column_letter(col)].width = max(float(width), 12)
            sheet.row_dimensions[1].height = 42
            for row_num, row in enumerate(source["rows"], 2):
                changed_columns = set(row["changed_columns"])
                for col, (saved, number_format) in enumerate(zip(row["values"], row["formats"]), 1):
                    cell = sheet.cell(row_num, col)
                    _write_cell(cell, saved)
                    cell.number_format = number_format
                    if row_num % 2 == 0:
                        cell.fill = PatternFill("solid", fgColor="F3F7FA")
                    if col in changed_columns:
                        cell.fill = PatternFill("solid", fgColor=CHANGED_FILL)
                        cell.border = Border(bottom=Side(style="thin", color=CHANGED_BORDER))
                sheet.row_dimensions[row_num].height = 24
            last_data_row = 1 + len(source["rows"])
            sheet.auto_filter.ref = f"A1:{end}{last_data_row}"
            sheet.freeze_panes = "D2"
            legend_row = last_data_row + 2
            sheet.merge_cells(f"A{legend_row}:{end}{legend_row}")
            legend = sheet.cell(legend_row, 1)
            legend.value = "Ô vàng: cập nhật trong lượt nhập này; ô còn lại: dữ liệu đã có trong BK."
            legend.font = Font(bold=True, color="7A4B00", size=11)
            legend.fill = PatternFill("solid", fgColor=CHANGED_FILL)
            legend.alignment = Alignment(vertical="center")
            sheet.row_dimensions[legend_row].height = 28
            sheet.print_title_rows = "1:1"
            sheet.sheet_properties.pageSetUpPr.fitToPage = True
            sheet.page_setup.fitToWidth = 1
        return book

    @staticmethod
    def _verify(path: Path, snapshot: Mapping[str, Any]) -> None:
        book = load_workbook(path, read_only=True, data_only=False)
        try:
            if book.sheetnames != [source["name"] for source in snapshot["sheets"]]:
                raise PostingExportError("Danh sách sheet file xuất không khớp lượt nhập BK.")
            for source in snapshot["sheets"]:
                sheet = book[source["name"]]
                if [cell.value for cell in sheet[1]][:len(source["headers"])] != source["headers"]:
                    raise PostingExportError("Header file xuất không khớp BK.")
                if sheet.max_row != len(source["rows"]) + 3:
                    raise PostingExportError("Số dòng file xuất không khớp lượt nhập BK.")
                for row_num, row in enumerate(source["rows"], 2):
                    for col, saved in enumerate(row["values"], 1):
                        cell = sheet.cell(row_num, col)
                        coordinate = f"{get_column_letter(col)}{row_num}"
                        if cell.value != _decoded_value(saved, coordinate):
                            raise PostingExportError("Giá trị trong file xuất không khớp BK.")
                        if col in row["changed_columns"] and cell.fill.fgColor.rgb[-6:] != CHANGED_FILL:
                            raise PostingExportError("Ô vừa cập nhật không được tô màu.")
        finally:
            book.close()
