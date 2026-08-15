from __future__ import annotations

import json
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook

from app.services.excel.carrier import (
    BK_DETAIL_SHEET,
    MISSING_INVOICE,
    PAYMENT_DETAIL_SHEET,
    PAYMENT_SUMMARY_SHEET,
    DETAIL_HEADERS,
    sync_payment_summary_sheet,
    upsert_bk_detail_rows,
)
from app.services.excel.models import ConflictType, ResolutionAction
from app.services.excel.payment_sync import SUMMARY_HEADERS, PaymentSyncService
from app.services.excel.posting import ExpensePostingService


class _Provider:
    def __init__(self, path: Path) -> None:
        self.path = path

    def get_ready_json_path(self, _batch_id: int) -> Path:
        return self.path

    def get_latest_ready_json_path(self) -> Path:
        return self.path


def _posting_book(path: Path, *, hp_carrier: str | None = None, nam_carrier: str | None = None) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "T07 26"
    headers = (
        "SQT PM",
        "Ngày Đóng",
        "Số Container",
        "Hạ Hàng",
        "Số HĐ",
        "VT bộ",
        "Nâng Hàng",
        "Số HĐ",
        "VT TRONG NAM",
    )
    for column, header in enumerate(headers, 1):
        sheet.cell(1, column).value = header
    sheet.append((700, "2026-07-01", "CONT700", None, None, hp_carrier, None, None, nam_carrier))
    workbook.save(path)
    workbook.close()


def _ready(path: Path, rows: list[list[object]]) -> None:
    path.write_text(json.dumps({"v": 1, "d": rows}, ensure_ascii=False), encoding="utf-8")


def test_posting_routes_expense_invoice_carriers_to_vt_nam(tmp_path: Path) -> None:
    bk = tmp_path / "bk.xlsx"
    ready = tmp_path / "ready.json"
    _posting_book(bk)
    _ready(
        ready,
        [
            ["CONT700", None, "HH", "CV", "HD-HP", "PHB", 1_200_000],
            ["CONT700", None, "NH", "CV", "HD-NAM", "NHS", 500_000],
        ],
    )
    service = ExpensePostingService(
        _Provider(ready),
        bk_path=bk,
        backup_dir=tmp_path / "Backup",
    )

    plan = service.analyze(batch_id=1, sheet_name="T07 26")
    assert not plan.conflicts
    result = service.apply(plan, {})

    assert result.carrier_written_cells == 1
    workbook = load_workbook(bk, data_only=False)
    try:
        assert workbook["T07 26"]["F2"].value is None
        assert workbook["T07 26"]["I2"].value == "PHB / NHS"
        assert workbook[BK_DETAIL_SHEET].sheet_state == "hidden"
        detail = workbook[BK_DETAIL_SHEET]
        assert detail.max_row == 3
        assert {detail["H2"].value, detail["H3"].value} == {"NAM"}
        assert {detail["L2"].value, detail["L3"].value} == {"PHB", "NHS"}
    finally:
        workbook.close()


def _add_road_fee_columns(path: Path) -> None:
    workbook = load_workbook(path)
    sheet = workbook["T07 26"]
    sheet["J1"] = "Cước bộ đóng hàng"
    sheet["K1"] = "Số HĐ"
    sheet["L1"] = "Cước VTN"
    sheet["M1"] = "Số HĐ"
    workbook.save(path)
    workbook.close()


def test_posting_blank_daily_sync_carrier_keeps_existing_value(
    tmp_path: Path,
) -> None:
    bk = tmp_path / "bk.xlsx"
    ready = tmp_path / "ready.json"
    _posting_book(bk, hp_carrier="DAILY ROAD")
    _ready(
        ready,
        [["CONT700", None, "CBDH", "CV", "HD-1", None, 100]],
    )
    _add_road_fee_columns(bk)
    service = ExpensePostingService(
        _Provider(ready), bk_path=bk, backup_dir=tmp_path / "Backup"
    )

    plan = service.analyze(batch_id=1, sheet_name="T07 26")
    assert not any(
        conflict.conflict_type is ConflictType.CARRIER_VALUE_CONFLICT
        for conflict in plan.conflicts
    )
    service.apply(plan, {})

    workbook = load_workbook(bk, data_only=False)
    try:
        assert workbook["T07 26"]["F2"].value == "DAILY ROAD"
    finally:
        workbook.close()


def test_posting_manual_daily_sync_carrier_overrides_existing_value(
    tmp_path: Path,
) -> None:
    bk = tmp_path / "bk.xlsx"
    ready = tmp_path / "ready.json"
    _posting_book(bk, hp_carrier="DAILY ROAD")
    _ready(
        ready,
        [["CONT700", None, "VTN", "CV", "HD-2", "USER ROAD", 200]],
    )
    _add_road_fee_columns(bk)
    service = ExpensePostingService(
        _Provider(ready), bk_path=bk, backup_dir=tmp_path / "Backup"
    )

    plan = service.analyze(batch_id=1, sheet_name="T07 26")
    assert not any(
        conflict.conflict_type is ConflictType.CARRIER_VALUE_CONFLICT
        for conflict in plan.conflicts
    )
    service.apply(plan, {})

    workbook = load_workbook(bk, data_only=False)
    try:
        assert workbook["T07 26"]["F2"].value == "USER ROAD"
    finally:
        workbook.close()


def test_posting_row_picker_candidates_include_json_and_bk_carriers(
    tmp_path: Path,
) -> None:
    bk = tmp_path / "bk.xlsx"
    ready = tmp_path / "ready.json"
    _posting_book(bk, hp_carrier="BK-HP", nam_carrier="BK-NAM")
    _ready(ready, [[None, "BL-01", "CB", "CV", "HD-01", "JSON-VT", 500_000]])
    service = ExpensePostingService(
        _Provider(ready), bk_path=bk, backup_dir=tmp_path / "Backup"
    )

    plan = service.analyze(batch_id=1, sheet_name="T07 26")
    conflict = next(
        value
        for value in plan.conflicts
        if value.conflict_type is ConflictType.BL_ONLY_NO_CONTAINER
    )

    assert conflict.carrier == "JSON-VT"
    assert [candidate.carrier for candidate in conflict.row_candidates] == [
        "BK-HP / BK-NAM"
    ]


def test_posting_carrier_conflict_append_only_for_different_invoice(tmp_path: Path) -> None:
    bk = tmp_path / "bk.xlsx"
    ready = tmp_path / "ready.json"
    _posting_book(bk, nam_carrier="NHS")
    workbook = load_workbook(bk)
    workbook["T07 26"]["H2"] = "HD-OLD"
    workbook.save(bk)
    workbook.close()
    _ready(ready, [["CONT700", None, "NH", "CV", "HD-NEW", "TP", 500_000]])
    service = ExpensePostingService(
        _Provider(ready), bk_path=bk, backup_dir=tmp_path / "Backup"
    )

    plan = service.analyze(batch_id=1, sheet_name="T07 26")
    conflict = next(
        value
        for value in plan.conflicts
        if value.conflict_type is ConflictType.CARRIER_VALUE_CONFLICT
    )
    assert ResolutionAction.APPEND_CARRIER in conflict.allowed_actions
    service.apply(
        plan,
        {conflict.conflict_id: {"action": "APPEND_CARRIER"}},
    )
    workbook = load_workbook(bk)
    try:
        assert workbook["T07 26"]["I2"].value == "NHS / TP"
    finally:
        workbook.close()


def test_posting_same_invoice_does_not_offer_append_carrier(tmp_path: Path) -> None:
    bk = tmp_path / "bk.xlsx"
    ready = tmp_path / "ready.json"
    _posting_book(bk, nam_carrier="NHS")
    workbook = load_workbook(bk)
    workbook["T07 26"]["H2"] = "HD-SAME"
    workbook.save(bk)
    workbook.close()
    _ready(ready, [["CONT700", None, "NH", "CV", "HD-SAME", "TP", 500_000]])
    service = ExpensePostingService(
        _Provider(ready), bk_path=bk, backup_dir=tmp_path / "Backup"
    )

    plan = service.analyze(batch_id=1, sheet_name="T07 26")
    conflict = next(
        value
        for value in plan.conflicts
        if value.conflict_type is ConflictType.CARRIER_VALUE_CONFLICT
    )

    assert set(conflict.allowed_actions) == {
        ResolutionAction.OVERWRITE,
        ResolutionAction.KEEP_EXISTING,
    }


def test_posting_keep_multiple_existing_carriers_requires_effective_choice(
    tmp_path: Path,
) -> None:
    bk = tmp_path / "bk.xlsx"
    ready = tmp_path / "ready.json"
    _posting_book(bk, nam_carrier="NHS / TP")
    workbook = load_workbook(bk)
    workbook["T07 26"]["H2"] = "HD-OLD"
    workbook.save(bk)
    workbook.close()
    _ready(ready, [["CONT700", None, "NH", "CV", "HD-NEW", "XYZ", 500_000]])
    service = ExpensePostingService(
        _Provider(ready), bk_path=bk, backup_dir=tmp_path / "Backup"
    )

    plan = service.analyze(batch_id=1, sheet_name="T07 26")
    conflict = next(
        value
        for value in plan.conflicts
        if value.conflict_type is ConflictType.CARRIER_VALUE_CONFLICT
    )
    assert conflict.details["requires_existing_carrier_selection"]
    service.apply(
        plan,
        {
            conflict.conflict_id: {
                "action": "KEEP_EXISTING",
                "selected_carrier": "TP",
            }
        },
    )

    workbook = load_workbook(bk, data_only=False)
    try:
        assert workbook["T07 26"]["I2"].value == "NHS / TP"
        assert workbook[BK_DETAIL_SHEET]["L2"].value == "TP"
    finally:
        workbook.close()


def test_selecting_ambiguous_invoice_carrier_keeps_other_invoice_carrier(
    tmp_path: Path,
) -> None:
    bk = tmp_path / "bk.xlsx"
    ready = tmp_path / "ready.json"
    _posting_book(bk)
    workbook = load_workbook(bk)
    sheet = workbook["T07 26"]
    for column, header in ((10, "Nâng vỏ"), (11, "Số HĐ"), (12, "Hạ vỏ"), (13, "Số HĐ")):
        sheet.cell(1, column).value = header
    workbook.save(bk)
    workbook.close()
    _ready(
        ready,
        [
            ["CONT700", None, "NH", "CV", "HD-A", "NHS", 500_000],
            ["CONT700", None, "NV", "CV", "HD-A", "TP", 600_000],
            ["CONT700", None, "HV", "CV", "HD-B", "XYZ", 700_000],
        ],
    )
    service = ExpensePostingService(
        _Provider(ready), bk_path=bk, backup_dir=tmp_path / "Backup"
    )

    plan = service.analyze(batch_id=1, sheet_name="T07 26")
    conflict = next(
        value
        for value in plan.conflicts
        if value.conflict_type is ConflictType.MULTIPLE_SOURCE_CARRIERS
    )
    refined = service.refine(
        plan,
        {
            conflict.conflict_id: {
                "action": "SELECT_CARRIER",
                "selected_carrier": "NHS",
            }
        },
    )
    assert not any(
        value.conflict_type is ConflictType.MULTIPLE_SOURCE_CARRIERS
        for value in refined.conflicts
    )
    service.apply(refined, {})

    workbook = load_workbook(bk, data_only=False)
    try:
        assert workbook["T07 26"]["I2"].value == "NHS / XYZ"
    finally:
        workbook.close()


def _payment_source(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "T07 26"
    headers = (
        "SQT", "Số Container", "Cước biển", "Cước bộ đóng hàng",
        "Nâng vỏ", "Số HĐ", "Hạ Hàng", "Số HĐ", "Nâng Hàng", "Số HĐ",
        "Hạ vỏ", "Số HĐ", "Cước VTN", "Lưu cont", "Số HĐ", "Quá tải",
        "Số HĐ", "VS + D/O", "Số HĐ", "LÀM LỆNH", "SỬA CHỮA", "Số HĐ",
        "VT bộ", "VT TRONG NAM", *SUMMARY_HEADERS,
    )
    for column, header in enumerate(headers, 1):
        sheet.cell(1, column).value = header
    sheet["A2"], sheet["B2"] = 700, "CONT700"
    sheet["G2"], sheet["H2"], sheet["W2"] = 1_200_000, "HD-HP", "PHB"
    sheet["I2"], sheet["J2"], sheet["X2"] = 500_000, "HD-NAM", "NHS"
    for column in range(25, 35):
        letter = sheet.cell(1, column).column_letter
        sheet.cell(2, column).value = f"=0+{letter}3"
    upsert_bk_detail_rows(
        workbook,
        [
            {
                "batch_id": 1,
                "batch_hash": "hash-1",
                "source_item_index": 0,
                "period": "T07/2026",
                "sheet_name": "T07 26",
                "sqt": 700,
                "container": "CONT700",
                "carrier_group": "HP",
                "fee": "HH",
                "invoice_no": "HD-HP",
                "carrier_source": "PHB",
                "carrier_effective": "PHB",
                "amount": 1_200_000,
                "carrier_action": "OVERWRITE",
                "updated_at": datetime(2026, 8, 7, 10, 0),
            },
            {
                "batch_id": 1,
                "batch_hash": "hash-1",
                "source_item_index": 1,
                "period": "T07/2026",
                "sheet_name": "T07 26",
                "sqt": 700,
                "container": "CONT700",
                "carrier_group": "NAM",
                "fee": "NH",
                "invoice_no": "HD-NAM",
                "carrier_source": "NHS",
                "carrier_effective": "NHS",
                "amount": 500_000,
                "carrier_action": "OVERWRITE",
                "updated_at": datetime(2026, 8, 7, 10, 0),
            },
        ],
    )
    workbook.save(path)
    workbook.close()


def _payment_target(path: Path) -> None:
    workbook = Workbook()
    workbook.remove(workbook.active)
    hp = workbook.create_sheet("T07 26 HP")
    for column, header in enumerate(
        ("QT", "SỐ CONT", "HẠ HÀNG", "Số HD", "Ghi chú", "Vận Tải", "Date cập nhật"), 1
    ):
        hp.cell(7, column).value = header
    hp["A12"], hp["C12"] = "TỔNG TIỀN", "=SUM(C8:C11)"
    nam = workbook.create_sheet("T07 26 NAM")
    nam_headers = (
        "QT", "SỐ CONT", "NÂNG VỎ", "Số HD", "NÂNG HÀNG", "Số HD",
        "HẠ VỎ", "Số HD", "VS + D/O", "Số HD", "LÀM LỆNH", "Lưu Cont",
        "Số HD", "Sửa chữa Cont", "Số HD", "QUÁ TẢI", "Số HD", "Vận Tải",
        "Date cập nhật",
    )
    for column, header in enumerate(nam_headers, 1):
        nam.cell(7, column).value = header
    nam["A12"], nam["E12"] = "TỔNG TIỀN", "=SUM(E8:E11)"
    workbook.save(path)
    workbook.close()


def test_payment_sync_writes_carriers_and_builds_summary(tmp_path: Path) -> None:
    bk, payment = tmp_path / "bk.xlsx", tmp_path / "payment.xlsx"
    _payment_source(bk)
    _payment_target(payment)
    service = PaymentSyncService(
        bk_path=bk,
        payment_path=payment,
        backup_dir=tmp_path / "Backup",
    )

    plan = service.analyze(source_sheet_name="T07 26")
    result = service.apply(plan, {})

    assert result.target_results["HP"].carrier_written_cells == 1
    assert result.target_results["NAM"].carrier_written_cells == 1
    workbook = load_workbook(payment, data_only=False)
    try:
        assert workbook["T07 26 HP"]["F8"].value == "PHB"
        assert workbook["T07 26 NAM"]["R8"].value == "NHS"
        assert PAYMENT_SUMMARY_SHEET in workbook.sheetnames
        summary = workbook[PAYMENT_SUMMARY_SHEET]
        assert {summary["B2"].value, summary["B3"].value} == {"NHS", "PHB"}
        assert {summary["F2"].value, summary["F3"].value} == {500_000, 1_200_000}
        assert summary["F2"].data_type == "n"
        assert {summary["J2"].value, summary["J3"].value} == {"NHS", "PHB"}
        assert {summary["U2"].value, summary["U3"].value} == {500_000, 1_200_000}
        assert PAYMENT_DETAIL_SHEET in workbook.sheetnames
        assert workbook[PAYMENT_DETAIL_SHEET].sheet_state == "hidden"
    finally:
        workbook.close()

    second_plan = service.analyze(source_sheet_name="T07 26")
    second_result = service.apply(second_plan, {})
    assert second_result.status.value == "NO_CHANGES"
    assert second_result.backup_path is None
    workbook = load_workbook(payment, data_only=False)
    try:
        detail = workbook[PAYMENT_DETAIL_SHEET]
        assert sum(
            detail.cell(row, 2).value == "hash-1"
            for row in range(2, detail.max_row + 1)
        ) == 2
    finally:
        workbook.close()


def test_payment_sync_appends_carrier_without_blocking_amount_sync(
    tmp_path: Path,
) -> None:
    bk, payment = tmp_path / "bk.xlsx", tmp_path / "payment.xlsx"
    _payment_source(bk)
    _payment_target(payment)
    workbook = load_workbook(payment)
    hp = workbook["T07 26 HP"]
    hp["A8"], hp["B8"], hp["F8"] = 700, "CONT700", "NHS"
    workbook.save(payment)
    workbook.close()
    service = PaymentSyncService(
        bk_path=bk,
        payment_path=payment,
        backup_dir=tmp_path / "Backup",
    )

    plan = service.analyze(source_sheet_name="T07 26")
    conflict = next(
        value
        for value in plan.conflicts
        if value.conflict_type is ConflictType.CARRIER_VALUE_CONFLICT
        and value.details.get("target_type") == "HP"
    )
    result = service.apply(
        plan,
        {conflict.conflict_id: {"action": "APPEND_CARRIER"}},
    )

    assert result.target_results["HP"].carrier_appended_cells == 1
    workbook = load_workbook(payment, data_only=False)
    try:
        hp = workbook["T07 26 HP"]
        assert hp["F8"].value == "NHS / PHB"
        assert hp["C8"].value == 1_200_000
    finally:
        workbook.close()


def _carrier_detail(
    *,
    batch_hash: str,
    source_index: int,
    period: str,
    sheet_name: str,
    carrier: str | None,
    invoice: str | None,
    fee: str,
    amount: int,
    container: str = "CONT700",
) -> dict[str, object]:
    return {
        "Batch ID": 1,
        "Batch hash": batch_hash,
        "Dòng JSON": source_index,
        "Kỳ": period,
        "Sheet BK": sheet_name,
        "SQT": 700,
        "Container": container,
        "Nhóm": "HP" if fee == "HH" else "NAM",
        "Mã phí": fee,
        "Số HĐ": invoice,
        "Carrier nguồn": carrier,
        "Carrier hiệu lực": carrier,
        "Số tiền": amount,
        "Hành động": "OVERWRITE",
        "Ngày cập nhật": datetime(2026, 8, 9, 10, 0),
    }


def test_carrier_summary_aggregates_invoices_then_month_and_writes_values(
    tmp_path: Path,
) -> None:
    target = tmp_path / "payment.xlsx"
    workbook = Workbook()
    workbook.active.title = "T07 26 HP"
    details = [
        _carrier_detail(
            batch_hash="a",
            source_index=0,
            period="T07/2026",
            sheet_name="T07 26",
            carrier="PHB",
            invoice=" HD-01 ",
            fee="HH",
            amount=100,
        ),
        _carrier_detail(
            batch_hash="a",
            source_index=1,
            period="T07/2026",
            sheet_name="T07 26",
            carrier="phb",
            invoice="hd-01",
            fee="NH",
            amount=200,
            container="CONT701",
        ),
        _carrier_detail(
            batch_hash="a",
            source_index=2,
            period="T07/2026",
            sheet_name="T07 26",
            carrier="PHB",
            invoice="HD-02",
            fee="LC",
            amount=300,
        ),
        _carrier_detail(
            batch_hash="a",
            source_index=3,
            period="T07/2026",
            sheet_name="T07 26",
            carrier="PHB",
            invoice=None,
            fee="QT",
            amount=50,
        ),
    ]

    report = sync_payment_summary_sheet(
        workbook, details, source_sheet="T07 26"
    )
    workbook.save(target)
    workbook.close()

    assert report.carrier_count == 1
    assert report.invoice_count == 2
    assert report.selected_period_total == 650
    assert report.missing_invoice_items == 1
    workbook = load_workbook(target, data_only=True)
    try:
        summary = workbook[PAYMENT_SUMMARY_SHEET]
        assert [summary.cell(2, column).value for column in range(1, 7)] == [
            "T07/2026",
            "PHB",
            2,
            600,
            50,
            650,
        ]
        invoice_rows = {
            summary.cell(row, 11).value: [
                summary.cell(row, column).value for column in range(12, 22)
            ]
            for row in range(2, 5)
        }
        assert invoice_rows["HD-01"] == [100, 0, 200, 0, 0, 0, 0, 0, 0, 300]
        assert invoice_rows["HD-02"][-1] == 300
        assert invoice_rows[MISSING_INVOICE][-1] == 50
    finally:
        workbook.close()


def test_carrier_summary_warns_cross_carrier_duplicates_and_unmapped() -> None:
    workbook = Workbook()
    workbook.active.title = "T07 26 HP"
    details = [
        _carrier_detail(
            batch_hash="a",
            source_index=0,
            period="T07/2026",
            sheet_name="T07 26",
            carrier="PHB",
            invoice="HD-X",
            fee="HH",
            amount=100,
        ),
        _carrier_detail(
            batch_hash="b",
            source_index=0,
            period="T07/2026",
            sheet_name="T07 26",
            carrier="PHB",
            invoice="HD-X",
            fee="HH",
            amount=100,
        ),
        _carrier_detail(
            batch_hash="c",
            source_index=0,
            period="T07/2026",
            sheet_name="T07 26",
            carrier="NHS",
            invoice="HD-X",
            fee="NH",
            amount=200,
        ),
        _carrier_detail(
            batch_hash="d",
            source_index=0,
            period="T07/2026",
            sheet_name="T07 26",
            carrier=None,
            invoice="HD-Z",
            fee="HV",
            amount=50,
        ),
    ]

    report = sync_payment_summary_sheet(
        workbook, details, source_sheet="T07 26"
    )

    assert report.selected_period_total == 450
    assert report.cross_carrier_invoices == 1
    assert report.suspected_duplicate_items == 1
    assert report.unmapped_carrier_items == 1
    summary = workbook[PAYMENT_SUMMARY_SHEET]
    warnings = " | ".join(
        str(summary.cell(row, 23).value or "")
        for row in range(2, summary.max_row + 1)
    )
    assert "nhiều bên vận tải" in warnings
    assert "nghi trùng" in warnings
    assert "CHƯA XÁC ĐỊNH" in warnings
    workbook.close()


def test_carrier_summary_accumulates_periods_and_replaces_selected_period() -> None:
    workbook = Workbook()
    workbook.active.title = "T06 26 HP"
    june = [
        _carrier_detail(
            batch_hash="june",
            source_index=0,
            period="T06/2026",
            sheet_name="T06 26",
            carrier="PHB",
            invoice="HD-06",
            fee="HH",
            amount=600,
        )
    ]
    july = [
        _carrier_detail(
            batch_hash="july",
            source_index=0,
            period="T07/2026",
            sheet_name="T07 26",
            carrier="NHS",
            invoice="HD-07",
            fee="NH",
            amount=700,
        )
    ]

    sync_payment_summary_sheet(workbook, june, source_sheet="T06 26")
    sync_payment_summary_sheet(workbook, july, source_sheet="T07 26")
    unchanged = sync_payment_summary_sheet(workbook, july, source_sheet="T07 26")
    assert not unchanged.changed
    summary = workbook[PAYMENT_SUMMARY_SHEET]
    assert [summary["A2"].value, summary["A3"].value] == ["T07/2026", "T06/2026"]

    cleared = sync_payment_summary_sheet(workbook, [], source_sheet="T06 26")
    assert cleared.changed
    assert summary["A2"].value == "T07/2026"
    assert summary["A3"].value is None
    detail = workbook[PAYMENT_DETAIL_SHEET]
    assert detail.max_row == 2
    assert detail["E2"].value == "T07 26"
    workbook.close()


def test_carrier_summary_migrates_legacy_detail_columns_without_loss() -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = PAYMENT_SUMMARY_SHEET
    for offset, header in enumerate(DETAIL_HEADERS, 16):
        worksheet.cell(1, offset).value = header
    legacy = _carrier_detail(
        batch_hash="legacy",
        source_index=0,
        period="T06/2026",
        sheet_name="T06 26",
        carrier="PHB",
        invoice="HD-OLD",
        fee="HH",
        amount=600,
    )
    for offset, header in enumerate(DETAIL_HEADERS, 16):
        worksheet.cell(2, offset).value = legacy[header]
    current = _carrier_detail(
        batch_hash="current",
        source_index=0,
        period="T07/2026",
        sheet_name="T07 26",
        carrier="NHS",
        invoice="HD-NEW",
        fee="NH",
        amount=700,
    )

    report = sync_payment_summary_sheet(
        workbook, [current], source_sheet="T07 26"
    )

    assert report.changed
    assert PAYMENT_DETAIL_SHEET in workbook.sheetnames
    detail = workbook[PAYMENT_DETAIL_SHEET]
    assert detail.sheet_state == "hidden"
    assert {detail["B2"].value, detail["B3"].value} == {"legacy", "current"}
    summary = workbook[PAYMENT_SUMMARY_SHEET]
    assert summary["P1"].value == "VSDL"
    assert {summary["A2"].value, summary["A3"].value} == {
        "T06/2026",
        "T07/2026",
    }
    workbook.close()
