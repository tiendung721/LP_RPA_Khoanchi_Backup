from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import Workbook

from app.services.excel.carrier_export import (
    CarrierExportError, CarrierExportService, normalize_carrier, normalize_carriers,
)
from app.services.excel.payment_sync import BK_HEADER_ALIASES, SUMMARY_HEADERS
from app.services.excel.workbook import workbook_fingerprint
from app.database import Database
from app.repositories.excel_run_repository import ExcelRunRepository


FIELDS = (
    "sea_freight", "north_freight", "empty_lift", "loaded_drop",
    "loaded_lift", "empty_drop", "south_freight", "storage", "overweight",
    "vs_do", "seal", "command_fee", "repair",
)


def _sheet(book: Workbook, name: str, rows: list[dict[str, object]]) -> None:
    sheet = book.create_sheet(name)
    sheet.cell(1, 1, "Ngày")  # Headers need not begin in column A.
    columns = {"sqt": 2, "container": 3}
    for field in ("sqt", "container"):
        sheet.cell(1, columns[field], BK_HEADER_ALIASES[field][0])
    next_column = 4
    invoice_columns = {}
    for field in FIELDS:
        columns[field] = next_column
        sheet.cell(1, next_column, BK_HEADER_ALIASES[field][0])
        next_column += 1
        if field not in {"sea_freight", "north_freight", "south_freight", "command_fee"}:
            invoice_columns[field] = next_column
            sheet.cell(1, next_column, "Số HĐ Seal" if field == "seal" else "Số HĐ")
            next_column += 1
    columns["carrier_nam"] = next_column
    sheet.cell(1, next_column, BK_HEADER_ALIASES["carrier_nam"][0])
    for column, header in enumerate(SUMMARY_HEADERS, next_column + 1):
        sheet.cell(1, column, header)
    for index, data in enumerate(rows, 2):
        sheet.cell(index, columns["sqt"], data.get("sqt", index))
        sheet.cell(index, columns["container"], data.get("container", f"CONT{index}"))
        sheet.cell(index, columns["carrier_nam"], data.get("carrier"))
        for field in FIELDS:
            if field in data:
                sheet.cell(index, columns[field], data[field])
            if f"invoice_{field}" in data and field in invoice_columns:
                sheet.cell(index, invoice_columns[field], data[f"invoice_{field}"])


def _bk(path: Path) -> None:
    book = Workbook()
    book.remove(book.active)
    _sheet(book, "T08 26", [
        {"sqt": 1, "container": "AA", "carrier": "NHS BK08", "empty_lift": "=100+50", "invoice_empty_lift": "00017"},
        {"sqt": 1, "container": "AA", "carrier": "NHS BK T08", "loaded_lift": 200},
        {"sqt": 2, "container": "BB", "carrier": "A. Quít", "repair": 500},
        {"sqt": 3, "container": "CC", "carrier": "Anh Quít", "repair": 700},
        {"sqt": 4, "container": "DD", "carrier": None, "overweight": 100},
        {"sqt": 5, "container": "EE", "carrier": "NHS / Anh Quít", "storage": 100},
        {"sqt": 6, "container": "FF", "carrier": "NHS BK08", "loaded_drop": 999},
    ])
    _sheet(book, "T09 26", [
        {"sqt": 1, "container": "AA", "carrier": "NHS BK09", "seal": 10, "vs_do": 20,
         "command_fee": 30, "storage": 40, "repair": 50, "overweight": 60, "empty_drop": 70},
    ])
    _sheet(book, "T10 26", [])
    book.save(path)
    book.close()


def test_normalize_only_explicit_month_suffix_and_distinct_names() -> None:
    assert normalize_carrier(" NHS   BK05 ") == "NHS"
    assert normalize_carrier("NHS BK T06 / NHS T06") == "NHS"
    assert normalize_carrier("A. Quít") == "A. Quít"
    assert normalize_carrier("Anh Quít") == "Anh Quít"
    assert normalize_carrier("A Quít 06") == "A Quít 06"
    assert normalize_carrier("NHS / Anh Quít") is None
    assert normalize_carrier("NHS BK07, NHS T07") == "NHS"
    assert normalize_carriers("NHS BK07, Anh Quít") == ("NHS", "Anh Quít")
    assert normalize_carriers("NHS + Anh Quít") == ("NHS", "Anh Quít")
    assert normalize_carriers("NHS; Anh Quít") == ("NHS", "Anh Quít")
    assert normalize_carrier("0") is None


def test_export_does_not_assign_handling_fees_from_vt_trong_nam(tmp_path: Path) -> None:
    bk = tmp_path / "bk.xlsx"
    _bk(bk)
    before = workbook_fingerprint(bk)
    service = CarrierExportService(bk_path=bk)

    preview = service.analyze(["T10 26", "T09 26", "T08 26"])

    assert preview.sheet_names == ("T08 26", "T09 26", "T10 26")
    assert preview.unresolved_count == 6
    assert preview.carriers == ()
    assert all(row.carrier_names == () for row in preview.rows)
    with pytest.raises(CarrierExportError, match="trong danh"):
        service.export(preview, "nhs", tmp_path)
    assert workbook_fingerprint(bk) == before


def test_saved_export_preview_keeps_unassigned_rows(tmp_path: Path) -> None:
    bk = tmp_path / "bk.xlsx"
    _bk(bk)
    database = Database(tmp_path / "history.db")
    try:
        service = CarrierExportService(
            bk_path=bk, run_repository=ExcelRunRepository(database)
        )
        preview = service.analyze(["T08 26"])
        assert preview.run_id is not None
        assert preview.unresolved_count == 5
        assert preview.carriers == ()
        resumed = service.analyze(["T08 26"], resume_run_id=preview.run_id)
        assert resumed.unresolved_count == 5
        assert resumed.carriers == ()
    finally:
        database.close()
