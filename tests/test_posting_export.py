from __future__ import annotations

from datetime import datetime
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill

from app.database import Database
from app.repositories.excel_run_repository import ExcelRunRepository
from app.repositories.expense_posting_repository import ExpensePostingRepository
from app.services.excel.models import ExcelRunStatus, PostingPlan, PostingResult, WorkbookFingerprint
from app.services.excel.payment_sync import SUMMARY_HEADERS
from app.services.excel.posting_export import (
    CHANGED_FILL,
    PostingExportError,
    PostingExportService,
    build_posting_export_snapshot,
)


def _bk(path: Path) -> None:
    book = Workbook()
    first = book.active
    first.title = "T07 26"
    second = book.create_sheet("T08 26")
    for sheet in (first, second):
        headers = [f"Cột {col}" for col in range(1, 45)]
        headers[0] = "SQT"
        headers[1] = "Ngày Đóng"
        headers[2] = "Số Container"
        headers[12] = "Cước biển"
        headers[14] = "Hóa đơn cước biển"
        headers[19] = "Nâng vỏ"
        headers[20] = "Số HĐ Nâng vỏ"
        headers[21] = "Hạ Hàng"
        headers[23] = "VT TRONG NAM"
        headers[24] = "Cước VTN"
        for col, value in enumerate(headers, 1):
            sheet.cell(1, col, value)
        sheet["M1"].fill = PatternFill("solid", fgColor="FFFF00")
        for col, value in enumerate(SUMMARY_HEADERS, 45):
            sheet.cell(1, col, value)
        sheet.column_dimensions["C"].width = 22
    first["A2"] = 701
    first["B2"] = datetime(2026, 7, 10)
    first["B2"].number_format = "dd/mm/yyyy"
    first["C2"] = "CONT-701"
    first["M2"] = 300
    first["O2"] = "=1+1"
    first["O2"].data_type = "s"
    first["V2"] = 50  # An older fee, visible but not highlighted.
    first["A3"] = 702
    first["C3"] = "CONT-702"
    second["A2"] = 801
    second["C2"] = "CONT-801"
    second["X2"] = "VT Nam"
    second["Y2"] = 200
    second["A3"] = 802
    second["C3"] = "CONT-802"
    second["E3"] = 2
    second["M3"] = "=E3*100"
    book.save(path)
    book.close()


def _plan(bk: Path, run_id: int | None) -> PostingPlan:
    return PostingPlan(
        batch_id=1,
        batch_path=bk.parent / "ready.json",
        batch_hash="hash",
        target_path=bk,
        target_fingerprint=WorkbookFingerprint(0, 0, "hash"),
        items=[], conflicts=[], sheet_candidates=[], run_id=run_id,
        source_rows=[{"source_item_index": 0}],
    )


def test_export_uses_bk_rows_and_saved_values_across_months(tmp_path: Path) -> None:
    bk = tmp_path / "BK.xlsx"
    _bk(bk)
    database = Database(tmp_path / "state.db")
    runs = ExcelRunRepository(database)
    postings = ExpensePostingRepository(database)
    run = runs.create_run(operation="EXPENSE_POSTING", target_path=bk)
    actions = [
        {"sheet_name": "T07 26", "target_row": 2, "target_column": 13, "amount_write": True},
        {"sheet_name": "T07 26", "target_row": 2, "invoice_target_column": 15, "invoice_write": True},
        {"sheet_name": "T08 26", "target_row": 2, "carrier_target_column": 24, "carrier_write": True},
        {"sheet_name": "T08 26", "target_row": 2, "target_column": 25, "amount_write": True},
        {"sheet_name": "T08 26", "target_row": 3, "target_column": 13, "amount_write": True},
        {"sheet_name": "T07 26", "target_row": 3, "target_column": 13, "amount_write": False},
    ]
    snapshot = build_posting_export_snapshot(
        _plan(bk, run.id), actions,
        PostingResult(status=ExcelRunStatus.SUCCEEDED, target_path=bk, run_id=run.id),
    )
    assert snapshot["version"] == 2
    assert snapshot["row_count"] == 3
    assert snapshot["sqt_count"] == 3
    assert [item["name"] for item in snapshot["sheets"]] == ["T07 26", "T08 26"]
    assert snapshot["sheets"][0]["rows"][0]["changed_columns"] == [13, 15]
    postings.save_export_snapshot(run.id, snapshot)
    runs.finish_run(run.id, status="SUCCEEDED")

    # The export must still show the values at posting time after BK is edited.
    changed_book = load_workbook(bk)
    changed_book["T07 26"]["M2"] = 999
    changed_book.save(bk)
    changed_book.close()
    reopened = ExpensePostingRepository(Database(tmp_path / "state.db"))
    exporter = PostingExportService(ExcelRunRepository(reopened.database), reopened)
    preview = exporter.latest_preview()
    assert preview is not None
    assert preview.row_count == 3
    assert preview.sqt_count == 3
    assert preview.sheet_count == 2
    assert preview.filename.startswith("BK_vua_nhap_T07-26_T08-26_3SQT_")
    target = tmp_path / preview.filename
    result = exporter.export(run.id, target)
    assert result.row_count == 3
    assert exporter.latest_preview().last_export_path == target
    with pytest.raises(FileExistsError):
        exporter.export(run.id, target)

    book = load_workbook(target, data_only=False)
    try:
        assert book.sheetnames == ["T07 26", "T08 26"]
        first = book["T07 26"]
        second = book["T08 26"]
        assert first.max_column == second.max_column == 44
        assert first["M1"].fill.fgColor.rgb[-6:] == "FFFF00"
        assert first["M2"].value == 300
        assert first["M2"].fill.fgColor.rgb[-6:] == CHANGED_FILL
        assert first["O2"].value == "=1+1"
        assert first["O2"].data_type == "s"
        assert first["O2"].fill.fgColor.rgb[-6:] == CHANGED_FILL
        assert first["V2"].value == 50
        assert first["V2"].fill.fgColor.rgb[-6:] != CHANGED_FILL
        assert first["B2"].value == datetime(2026, 7, 10)
        assert first["B2"].number_format == "dd/mm/yyyy"
        assert first["A3"].value is None  # Unchanged SQT 702 is absent.
        assert second["A3"].value == 802
        assert second["M3"].value == "=E3*100"  # Original row 3 stays row 3 here.
        assert first.auto_filter.ref == "A1:AR2"
        assert first["A4"].value.startswith("Ô vàng:")
    finally:
        book.close()


def test_formula_rebases_when_original_bk_row_moves(tmp_path: Path) -> None:
    bk = tmp_path / "BK.xlsx"
    _bk(bk)
    snapshot = build_posting_export_snapshot(
        _plan(bk, 1),
        [{"sheet_name": "T08 26", "target_row": 3, "target_column": 13, "amount_write": True}],
        PostingResult(status=ExcelRunStatus.SUCCEEDED, target_path=bk, run_id=1),
    )
    book = PostingExportService._workbook(snapshot)
    try:
        assert book["T08 26"]["M2"].value == "=E2*100"
        assert book["T08 26"]["E2"].value == 2
    finally:
        book.close()


def test_latest_run_with_no_bk_changes_does_not_export_old_run(tmp_path: Path) -> None:
    bk = tmp_path / "BK.xlsx"
    _bk(bk)
    database = Database(tmp_path / "state.db")
    runs = ExcelRunRepository(database)
    postings = ExpensePostingRepository(database)
    first = runs.create_run(operation="EXPENSE_POSTING", target_path=bk)
    postings.save_export_snapshot(first.id, build_posting_export_snapshot(
        _plan(bk, first.id),
        [{"sheet_name": "T07 26", "target_row": 2, "target_column": 13, "amount_write": True}],
        PostingResult(status=ExcelRunStatus.SUCCEEDED, target_path=bk, run_id=first.id),
    ))
    runs.finish_run(first.id, status="SUCCEEDED")
    second = runs.create_run(operation="EXPENSE_POSTING", target_path=bk)
    postings.save_export_snapshot(second.id, build_posting_export_snapshot(
        _plan(bk, second.id), [],
        PostingResult(status=ExcelRunStatus.NO_CHANGES, target_path=bk, run_id=second.id),
    ))
    runs.finish_run(second.id, status="NO_CHANGES")
    exporter = PostingExportService(runs, postings)
    preview = exporter.latest_preview()
    assert preview is not None and preview.run_id == second.id and preview.row_count == 0
    with pytest.raises(PostingExportError, match="không có dòng BK"):
        exporter.export(second.id, tmp_path / "empty.xlsx")
    assert not (tmp_path / "empty.xlsx").exists()


def test_legacy_audit_snapshot_is_not_exposed_as_bk_format(tmp_path: Path) -> None:
    database = Database(tmp_path / "state.db")
    runs = ExcelRunRepository(database)
    postings = ExpensePostingRepository(database)
    run = runs.create_run(operation="EXPENSE_POSTING")
    postings.save_export_snapshot(run.id, {"version": 1, "rows": [{"status": "Đã ghi"}]})
    runs.finish_run(run.id, status="SUCCEEDED")
    exporter = PostingExportService(runs, postings)
    assert exporter.latest_preview() is None
    with pytest.raises(PostingExportError, match="format mới"):
        exporter.export(run.id, tmp_path / "old.xlsx")
